#!/usr/bin/env python3

import argparse
import sys
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:
    print("FAIL: openpyxl is required. Install with: python -m pip install openpyxl")
    sys.exit(2)


CLIENT_REQUIRED = {
    "脚本审核确认表": ["序号", "对应镜号/范围", "问题类型", "建议", "用户意见/选择"],
    "形象场景描述确认表": ["类别", "名称", "对应镜号/范围", "初步设定方向", "用户确认"],
}

INTERNAL_ASSET_REQUIRED = {
    "设定资产制作表": ["资产类别", "资产名称", "对应镜号/范围", "Prompt类型", "完整Prompt", "输出文件名", "制作状态"],
}

INTERNAL_PROMPT_REQUIRED = {
    "内部执行脚本与Prompt表": ["镜号", "原始脚本内容", "执行版画面内容", "首帧Prompt", "生成前执行说明", "视频内容Prompt（含声音/负面）"],
}


def normalized_headers(sheet):
    return [str(cell.value).strip() for cell in sheet[1] if cell.value is not None]


def find_sheet(workbook, expected_name):
    if expected_name in workbook.sheetnames:
        return workbook[expected_name]
    for name in workbook.sheetnames:
        if expected_name in name:
            return workbook[name]
    return None


def validate_required(workbook, required):
    failures = []
    for sheet_name, headers in required.items():
        sheet = find_sheet(workbook, sheet_name)
        if sheet is None:
            failures.append(f"missing sheet: {sheet_name}")
            continue
        actual = normalized_headers(sheet)
        for header in headers:
            if header not in actual:
                failures.append(f"{sheet.title} missing column: {header}")
        if sheet.max_row < 2:
            failures.append(f"{sheet.title} has no body rows")
    return failures


def main():
    parser = argparse.ArgumentParser(description="Validate 韦斯安德森 workbook structure.")
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--type", choices=["client", "internal-asset", "internal-prompt"], default="client")
    args = parser.parse_args()

    if not args.workbook.exists():
        print(f"FAIL: workbook does not exist: {args.workbook}")
        return 1

    workbook = load_workbook(args.workbook, read_only=True, data_only=False)

    if args.type == "client":
        failures = validate_required(workbook, CLIENT_REQUIRED)
        if workbook.sheetnames[:2] != ["脚本审核确认表", "形象场景描述确认表"]:
            failures.append("first two sheets must be 脚本审核确认表 and 形象场景描述确认表")
    elif args.type == "internal-asset":
        failures = validate_required(workbook, INTERNAL_ASSET_REQUIRED)
    else:
        failures = validate_required(workbook, INTERNAL_PROMPT_REQUIRED)

    if failures:
        for failure in failures:
            print(f"FAIL: {failure}")
        return 1

    print(f"PASS: {args.workbook} matches {args.type} workbook structure")
    return 0


if __name__ == "__main__":
    sys.exit(main())

