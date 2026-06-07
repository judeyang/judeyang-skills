#!/usr/bin/env python3
"""Build an internal asset production workbook scaffold from an audit workbook."""

from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


CENTER_HEADERS = {
    "序号",
    "类别",
    "项目",
    "对应镜号/范围",
    "Prompt类型",
    "素材状态",
    "确认状态",
    "用户确认",
    "是否发生在已确认之后",
}


def safe_text(value) -> str:
    return "" if value is None else str(value).strip()


def body_alignment(header_value):
    header = safe_text(header_value)
    if header in CENTER_HEADERS or header.endswith("编号"):
        return Alignment(horizontal="center", vertical="top", wrap_text=True)
    return Alignment(horizontal="left", vertical="top", wrap_text=True)


def style_sheet(ws, widths: list[int], row_height: int = 72):
    header_fill = PatternFill("solid", fgColor="1F4E78")
    thin = Side(style="thin", color="D9E2F3")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = Font(name="Arial", color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial", size=10)
            cell.alignment = body_alignment(ws.cell(1, cell.column).value)
            cell.border = border

    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = width
    for row_index in range(1, ws.max_row + 1):
        ws.row_dimensions[row_index].height = row_height if row_index > 1 else 30
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def find_assets_sheet(workbook: Path):
    wb = load_workbook(workbook, data_only=True)
    for name in ["形象场景描述确认表", "设定资产提取表", "资产提取表"]:
        if name in wb.sheetnames:
            return wb[name]
    raise RuntimeError("未找到形象场景描述确认表，需要先生成脚本审核确认表或资产提取表。")


def read_asset_rows(workbook: Path) -> list[dict[str, str]]:
    ws = find_assets_sheet(workbook)
    headers = {safe_text(cell.value): cell.column for cell in ws[1]}
    category_col = headers.get("类别")
    name_col = headers.get("项目")
    direction_col = headers.get("初步设定方向")
    confirm_col = headers.get("需要确认的问题")
    material_col = headers.get("是否需要客户提供参考")
    shot_col = headers.get("对应镜号/范围") or headers.get("对应镜号")
    source_col = headers.get("原始内容摘要") or headers.get("来源脚本摘要")

    if not category_col or not name_col:
        raise RuntimeError("形象场景描述确认表缺少必要列：类别、项目。")

    rows: list[dict[str, str]] = []
    for row_idx in range(2, ws.max_row + 1):
        category = safe_text(ws.cell(row_idx, category_col).value)
        name = safe_text(ws.cell(row_idx, name_col).value)
        if not category or not name:
            continue
        rows.append(
            {
                "category": category,
                "name": name,
                "shots": safe_text(ws.cell(row_idx, shot_col).value) if shot_col else "",
                "source": safe_text(ws.cell(row_idx, source_col).value) if source_col else "",
                "direction": safe_text(ws.cell(row_idx, direction_col).value) if direction_col else "",
                "material": safe_text(ws.cell(row_idx, material_col).value) if material_col else "",
                "question": safe_text(ws.cell(row_idx, confirm_col).value) if confirm_col else "",
            }
        )
    return rows


def prompt_types_for(category: str) -> list[str]:
    if "人物" in category:
        return ["三视图Prompt", "大头照Prompt", "全身照Prompt"]
    if "动物" in category or "坐骑" in category:
        return ["三视图Prompt", "大头照Prompt", "全身照Prompt"]
    if "场景" in category:
        return ["场景设定Prompt"]
    if "产品" in category or "品牌" in category:
        return ["产品参考Prompt"]
    if "道具" in category or "法器" in category:
        return ["道具设定Prompt"]
    if "特效" in category:
        return ["特效关键帧Prompt"]
    if "字体" in category or "包装" in category or "文字" in category:
        return ["字体包装参考Prompt"]
    return ["设定资产Prompt"]


def default_output_path(input_path: Path, project_dir: Path | None) -> Path:
    if project_dir:
        target_dir = project_dir / "05_内部制作执行"
        target_dir.mkdir(parents=True, exist_ok=True)
    else:
        target_dir = input_path.parent

    for version in range(1, 100):
        candidate = target_dir / f"设定资产制作表_v{version:02d}.xlsx"
        if not candidate.exists():
            return candidate
    raise RuntimeError("无法生成输出文件名：设定资产制作表_v01-v99 已存在。")


def write_workbook(asset_rows: list[dict[str, str]], output_path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "设定资产制作表"
    ws.append(
        [
            "序号",
            "类别",
            "项目",
            "对应镜号/范围",
            "来源脚本摘要",
            "初步设定方向",
            "Prompt类型",
            "完整Prompt",
            "生成输出文件名",
            "成品资产路径",
            "素材状态",
            "制作状态",
            "制作备注",
            "备注",
        ]
    )

    index = 1
    for row in asset_rows:
        material = row["material"] or "如客户有指定参考图/品牌素材请提供；没有则按确认后的文字方向生成。"
        material_status = "待提供" if any(key in material for key in ["需要", "请提供", "官方", "Logo", "品牌"]) else "待确认"
        for prompt_type in prompt_types_for(row["category"]):
            ws.append(
                [
                    index,
                    row["category"],
                    row["name"],
                    row["shots"],
                    row["source"],
                    row["direction"],
                    prompt_type,
                    "",
                    "",
                    "",
                    material_status,
                    "待制作",
                    "",
                    "本表为内部资产制作表。Prompt用于制作人员出图；完成后把实际图片文件名和确认项更新回项目名_脚本与资产确认表_v01.xlsx。",
                ]
            )
            index += 1

    ws_feedback = wb.create_sheet("用户修改意见提交模板")
    ws_feedback.append(
        [
            "反馈时间",
            "反馈来源",
            "关联文件",
            "关联镜头/资产",
            "客户原话",
            "要求类型",
            "是否发生在已确认之后",
            "是否已有样片或图片受影响",
            "需要修改的内容",
            "不能修改的内容",
            "客户提供的新素材",
            "期望完成时间",
            "我的处理建议/备注",
            "处理状态",
            "下一步动作",
        ]
    )
    ws_feedback.append(["", "微信 / 飞书 / 邮件 / 电话 / 会议 / 文件批注 / 其他", "", "", "", "脚本修改 / 人物设定 / 场景设定 / 产品素材 / 道具特效 / 配音音乐 / 字幕包装 / 成片节奏 / 其他", "是 / 否 / 不确定", "是 / 否 / 不确定", "", "", "", "", "", "待处理", ""])

    style_sheet(ws, [8, 14, 20, 16, 34, 38, 18, 54, 26, 34, 14, 16, 32, 44], row_height=84)
    style_sheet(ws_feedback, [18, 18, 24, 20, 44, 20, 20, 24, 36, 30, 28, 18, 34, 16, 30], row_height=70)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True, help="Audit workbook containing 形象场景描述确认表")
    parser.add_argument("--output", help="Output internal asset production workbook path")
    parser.add_argument("--project-dir", help="Optional project root folder")
    args = parser.parse_args()

    input_path = Path(args.input).expanduser()
    project_dir = Path(args.project_dir).expanduser() if args.project_dir else None
    output_path = Path(args.output).expanduser() if args.output else default_output_path(input_path, project_dir)
    asset_rows = read_asset_rows(input_path)
    write_workbook(asset_rows, output_path)
    print(output_path)


if __name__ == "__main__":
    main()
