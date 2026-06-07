#!/usr/bin/env python3
"""Build an internal execution workbook scaffold.

This script intentionally does not generate any prompt content. Prompt cells are
left blank for the AI director pass after the full script has been read.
"""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

CENTER_HEADERS = {
    "序号",
    "镜号",
    "原始镜号",
    "时长",
    "执行时长(s)",
    "景别",
    "制作状态",
    "输出文件名",
}

COMPACT_TEXT_HEADERS = {
    "确认后台词/旁白",
    "台词/旁白",
    "音效/音乐",
    "样片建议",
}


def body_alignment(header_value):
    header = str(header_value or "").strip()
    if header in CENTER_HEADERS or header.endswith("编号"):
        return Alignment(horizontal="center", vertical="top", wrap_text=True)
    if header in COMPACT_TEXT_HEADERS:
        return Alignment(horizontal="left", vertical="top", wrap_text=True)
    return Alignment(horizontal="left", vertical="top", wrap_text=True)


IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp"}


def normalize(value) -> str:
    return str(value or "").strip().replace("\n", " / ")


def find_project_root(path: Path) -> Path | None:
    for parent in [path.parent, *path.parents]:
        if (
            parent.name.endswith("_项目交付文件夹")
            or (parent / "03_设定资产").exists()
            or (parent / "04_设定资产").exists()
            or (parent / "03_设定图输出").exists()
        ):
            return parent
    return None


def find_header_row(ws) -> tuple[int, dict[str, int]] | None:
    canonical = {
        "shot": ["镜号", "执行镜号"],
        "orig": ["原始镜号", "对应原始镜号"],
        "orig_text": ["原始脚本内容"],
        "picture": ["确认后执行内容", "确认后执行画面", "执行版画面内容", "画面内容", "画面"],
        "dialogue": ["确认后台词/旁白", "执行版台词/旁白", "台词/旁白", "台词"],
        "duration": ["时长", "执行时长"],
        "sfx": ["音效/音乐", "音效", "音乐"],
        "note": ["改动说明", "执行备注", "备注"],
    }
    for row_idx in range(1, min(ws.max_row, 30) + 1):
        values = [normalize(ws.cell(row_idx, col).value) for col in range(1, ws.max_column + 1)]
        mapping: dict[str, int] = {}
        for col_idx, value in enumerate(values, 1):
            compact = value.replace(" ", "")
            for key, names in canonical.items():
                if key not in mapping and any(name in compact for name in names):
                    mapping[key] = col_idx
        if "shot" in mapping and "picture" in mapping:
            return row_idx, mapping
    return None


def read_execution_rows(path: Path) -> list[dict[str, str]]:
    wb = load_workbook(path, data_only=True)
    candidates = [wb["执行版脚本"]] if "执行版脚本" in wb.sheetnames else []
    candidates.extend(ws for ws in wb.worksheets if ws not in candidates)
    ws = None
    found = None
    for candidate in candidates:
        found = find_header_row(candidate)
        if found:
            ws = candidate
            break
    if not found:
        raise RuntimeError("未找到执行脚本表头，需要至少包含 镜号 和 确认后执行内容/确认后执行画面。")
    assert ws is not None
    header_row, mapping = found
    rows: list[dict[str, str]] = []
    for row_idx in range(header_row + 1, ws.max_row + 1):
        shot = normalize(ws.cell(row_idx, mapping["shot"]).value)
        if not shot:
            continue
        if shot in {"合计", "总计"}:
            rows.append({"total": "1", "duration": normalize(ws.cell(row_idx, mapping.get("duration", 0)).value) if mapping.get("duration") else ""})
            continue
        rows.append(
            {
                "shot": shot,
                "orig": normalize(ws.cell(row_idx, mapping.get("orig", 0)).value) if mapping.get("orig") else "",
                "orig_text": normalize(ws.cell(row_idx, mapping.get("orig_text", 0)).value) if mapping.get("orig_text") else "",
                "picture": normalize(ws.cell(row_idx, mapping["picture"]).value),
                "dialogue": normalize(ws.cell(row_idx, mapping.get("dialogue", 0)).value) if mapping.get("dialogue") else "",
                "duration": normalize(ws.cell(row_idx, mapping.get("duration", 0)).value) if mapping.get("duration") else "",
                "sfx": normalize(ws.cell(row_idx, mapping.get("sfx", 0)).value) if mapping.get("sfx") else "",
                "note": normalize(ws.cell(row_idx, mapping.get("note", 0)).value) if mapping.get("note") else "",
            }
        )
    return rows


def scan_assets(project_dir: Path | None) -> dict[str, Path]:
    if not project_dir:
        return {}
    root = project_dir / "03_设定资产"
    if not root.exists():
        root = project_dir / "04_设定资产"
    if not root.exists():
        root = project_dir / "03_设定图输出"
    if not root.exists():
        return {}
    assets: dict[str, Path] = {}
    for path in root.rglob("*"):
        if path.is_file() and path.suffix.lower() in IMAGE_EXTENSIONS:
            assets[path.stem] = path
    return assets


def asset_display(path: Path) -> str:
    return path.name


def infer_project_style(rows: list[dict[str, str]]) -> str:
    text = " ".join(f"{row.get('orig_text','')} {row.get('picture','')} {row.get('dialogue','')}" for row in rows)
    if any(k in text for k in ["红楼", "大观园", "贾母", "黛玉", "宝钗", "宝玉", "王熙凤", "刘姥姥"]):
        return "写实电影感，古典红楼梦气质，80年代古装影视质感参考但不复刻具体影视版本；人物、服装、妆发和场景构图保持原创。大观园宴会厅明亮雅致，木质梁柱、屏风、帷幔、烛光、瓷器材质清晰。柔和暖光，轻微胶片颗粒，真实皮肤和织物纹理。"
    if any(k in text for k in ["西游", "悟空", "八戒", "唐僧", "沙僧", "黄风怪", "花果山"]):
        return "经典神话剧质感，复古棚拍与电影级光影结合，真实材质，统一人物造型和服化道；参考传统古装神话美学但不复刻具体影视版本。画面干净高级，适合品牌广告短片。"
    return "写实电影感，统一人物造型、场景空间和品牌质感；光线、色彩、镜头语言服务剧情情绪，主体清晰可辨。"


def infer_continuity(rows: list[dict[str, str]]) -> str:
    text = " ".join(f"{row.get('orig_text','')} {row.get('picture','')} {row.get('dialogue','')}" for row in rows)
    rules: list[str] = []
    if any(k in text for k in ["大观园", "贾母", "黛玉", "宝钗", "宝玉", "空座", "C位"]):
        rules.extend(
            [
                "大观园宴会厅沿用统一空间轴线：镜头面向贾母主位，贾母居中靠后；贾母右手边的锦绣空座是本段争夺的C位，不得在镜头之间漂移。",
                "从观众视角锁定三人关系：林黛玉在画面左侧/空座左侧，薛宝钗在画面右侧/空座右侧，贾宝玉从两人附近起身调停，后续三人对峙保持宝玉居中、黛玉左、宝钗右。",
                "同一宴席段落中桌面、屏风、烛火、红绸、空座和人物左右关系保持一致；除非脚本明确写横移或反打，不得随意翻转左右方向。",
            ]
        )
    if any(k in text for k in ["空调", "产品", "出风", "风口", "屏显", "启动"]):
        rules.append("产品出现后保持同一产品位置和状态逻辑：未启动、启动、屏显、出风口/导风板开启等变化必须按脚本顺序推进，不得在前一镜头未启动时提前出现启动后的状态。")
    if any(k in text for k in ["西游", "黄风岭", "花果山", "黄风怪", "灵吉菩萨"]):
        rules.append("西游段落保持同一取经队伍方向和场景变化逻辑：黄风岭到雨林/花果山的转变由风或产品效果推动，角色站位和风向在连续镜头中保持可追踪。")
    if not rules:
        rules.append("全片连续性需先按完整脚本锁定：同一场景的人物站位、道具位置、产品状态、光线方向和镜头轴线在相邻镜头中保持一致；任何位置变化必须由脚本动作解释。")
    return "\n".join(f"- {rule}" for rule in rules)


ASSET_ALIASES = {
    "贾母": ["贾母"],
    "林黛玉": ["林黛玉", "黛玉"],
    "薛宝钗": ["薛宝钗", "宝钗"],
    "贾宝玉": ["贾宝玉", "宝玉"],
    "王熙凤": ["王熙凤", "凤姐"],
    "刘姥姥": ["刘姥姥"],
    "群像": ["群像", "众人", "宴会", "宴席"],
    "宴会厅": ["宴会厅", "大观园", "宴席"],
    "C位座次": ["C位", "空座", "座次"],
    "道具": ["团扇", "绢帕", "器物", "杯盏", "道具"],
    "气流": ["气流", "送风", "170度", "环抱风"],
    "孙悟空": ["孙悟空", "悟空"],
    "唐三藏": ["唐三藏", "唐僧", "师父"],
    "猪八戒": ["猪八戒", "八戒", "老猪"],
    "沙僧": ["沙僧"],
    "白龙马": ["白龙马", "马"],
    "黄风怪": ["黄风怪"],
    "灵吉菩萨": ["灵吉菩萨", "菩萨"],
    "黄风岭": ["黄风岭", "黄沙", "风沙", "沙丘"],
    "花果山": ["花果山", "雨林", "瀑布", "彩虹"],
    "海尔空调": ["海尔", "空调", "产品", "水洗空气"],
    "卡萨帝空调": ["卡萨帝", "空调", "产品", "柜机", "揽光"],
    "法器": ["净风宝瓶", "定风珠", "法器", "神器"],
}


def matching_asset_paths(label: str, keywords: list[str], assets: dict[str, Path]) -> list[Path]:
    character_labels = {"贾母", "林黛玉", "薛宝钗", "贾宝玉", "王熙凤", "刘姥姥", "孙悟空", "唐三藏", "猪八戒", "沙僧", "白龙马", "黄风怪", "灵吉菩萨"}
    scene_labels = {"宴会厅", "C位座次", "黄风岭", "花果山"}
    matches: list[tuple[int, Path]] = []
    for name, path in assets.items():
        searchable = f"{name} {path.name}"
        if label in searchable or any(k in searchable for k in keywords):
            score = 5
            if label in searchable:
                score = 2
            if label in character_labels and f"角色_{label}" in searchable:
                score = 0
            if label in scene_labels and ("场景_" in searchable or label in searchable):
                score = min(score, 1)
            if label == "群像" and "群像" in searchable:
                score = 0
            if label == "道具" and "道具" in searchable:
                score = 0
            matches.append((score, path))
    return [path for _, path in sorted(matches, key=lambda item: (item[0], len(item[1].name), item[1].name))]


def references_for(row: dict[str, str], assets: dict[str, Path]) -> str:
    dialogue = row.get("dialogue", "")
    if "标题字" in dialogue:
        dialogue = ""
    # Use confirmed execution content for visual references. Original script text may
    # contain titles or replaced ideas that are no longer visible in the shot.
    text = f"{row.get('picture','')} {dialogue}"
    lines: list[str] = []
    for label, keywords in ASSET_ALIASES.items():
        if not any(k in text for k in keywords):
            continue
        matches = matching_asset_paths(label, keywords, assets)
        if matches:
            lines.append(f"{label}: {asset_display(matches[0])}")
        else:
            lines.append(f"{label}: 待补充")
    return "\n".join(lines) if lines else "按已确认设定和文字方向执行"


def shot_style(row: dict[str, str]) -> tuple[str, str, str]:
    text = row.get("picture", "")
    view = "中景"
    if "特写" in text:
        view = "特写或近景"
    elif "近景" in text:
        view = "近景"
    elif "全景" in text:
        view = "全景"

    if any(k in text for k in ["黄风岭", "黄沙", "风沙", "沙丘"]):
        light = "昏黄风沙色调，低饱和黄褐色，空气中有细密沙尘，暗部保留细节。"
    elif any(k in text for k in ["花果山", "雨林", "瀑布", "彩虹", "绿叶"]):
        light = "雨后清新绿色调，空气通透，水汽柔和，画面干净明亮。"
    elif any(k in text for k in ["空调", "产品", "Logo"]):
        light = "品牌广告质感，产品边缘干净高光，机身结构清晰，Logo不可错误。"
    else:
        light = "电影级真实光影，色彩服务剧情情绪，主体清晰可辨。"

    subject = subject_hint(text)
    camera = (
        f"景别：{view}。机位：平视机位。"
        f"构图：{subject}占据竖屏主要阅读区，动作方向一侧留出可见余量，背景只保留能说明场景关系的元素。"
        f"运镜：以{subject}的动作节奏为准，简单反应镜头保持稳定，动作推进时做克制跟随或轻推。"
    )
    first_frame = f"{row.get('picture', '')} 首帧需要明确主体位置、场景环境、动作起点和产品/角色可见状态。"
    return camera, first_frame, light


def subject_hint(text: str) -> str:
    candidates = [
        "贾母", "林黛玉", "薛宝钗", "贾宝玉", "王熙凤", "刘姥姥",
        "孙悟空", "唐三藏", "猪八戒", "沙僧", "黄风怪", "灵吉菩萨",
        "空调", "产品", "C位空座", "宴席众人", "气流", "风"
    ]
    found = [name for name in candidates if name in text]
    if found:
        return "、".join(found[:3])
    return "本镜头主要人物或产品"


def shot_core_theme(row: dict[str, str]) -> str:
    text = row.get("picture", "")
    dialogue = row.get("dialogue", "")
    subject = subject_hint(f"{text} {dialogue}")
    if movement_keywords(text):
        return f"{subject}的状态变化或产品效果展示，重点是动作路径、起止状态和视觉因果清楚。"
    if dialogue:
        return f"{subject}的台词/反应表演，重点是表情、视线、手部动作和人物关系连续。"
    return f"{subject}的剧情推进镜头，重点是空间关系、动作起点和结尾状态清楚。"


def motion_rule_for(row: dict[str, str]) -> str:
    text = row.get("picture", "")
    duration = parse_duration_seconds(row.get("duration", "")) or 0
    subject = subject_hint(text)
    if movement_keywords(text):
        return (
            f"本镜头按连续动作处理，镜头先锁住{subject}的起始状态，再随变化路径克制推进；"
            "不得突然跳到完成状态，结尾必须留出稳定停顿。"
        )
    if duration >= 5:
        return (
            f"本镜头按长台词表演处理，镜头服务{subject}的表情和视线变化；"
            "允许轻微推近或切到中近景，但不得无理由跳轴。"
        )
    return (
        f"本镜头按短反应/过渡镜头处理，镜头稳定呈现{subject}的动作起点和结束姿态；"
        "不额外增加花哨运动。"
    )


def parse_duration_seconds(value: str) -> float | None:
    text = normalize(value)
    if not text:
        return None
    match = re.search(r"(\d+(?:\.\d+)?)", text)
    if not match:
        return None
    try:
        return float(match.group(1))
    except ValueError:
        return None


MUSIC_MARKERS = ["音乐", "配乐", "背景乐", "BGM", "bgm", "古筝", "琵琶", "合奏", "变奏", "轮指", "弦乐", "鼓点", "旋律"]


def video_sound_effects(value: str) -> str:
    text = normalize(value)
    if not text:
        return "只保留现场环境声和必要动作音效"
    parts = [part.strip(" ，,、；;") for part in re.split(r"[，,、；;]", text) if part.strip(" ，,、；;")]
    effects = [part for part in parts if not any(marker in part for marker in MUSIC_MARKERS)]
    return "，".join(effects) if effects else "只保留现场环境声和必要动作音效"


def format_time(seconds: float) -> str:
    if abs(seconds - round(seconds)) < 0.01:
        return f"00:{int(round(seconds)):02d}"
    whole = int(seconds)
    decimal = int(round((seconds - whole) * 10))
    return f"00:{whole:02d}.{decimal}"


def split_ranges(duration: float | None, preferred_segments: int) -> list[tuple[float, float]]:
    if not duration or duration <= 0:
        return [(0, 0)]
    if duration <= 2:
        return [(0, duration)]
    segment_count = max(2, min(preferred_segments, 4))
    if segment_count == 3:
        if duration <= 4:
            cuts = [round(duration / 3, 1), round(duration * 2 / 3, 1)]
        elif duration <= 5:
            cuts = [1.5, 3.0]
        else:
            cuts = [2.0, 4.0]
        return [(0, cuts[0]), (cuts[0], cuts[1]), (cuts[1], duration)]
    if segment_count == 2:
        cut = round(duration / 2, 1) if duration <= 4 else 2.0
        return [(0, cut), (cut, duration)]
    step = duration / segment_count
    ranges: list[tuple[float, float]] = []
    start = 0.0
    for index in range(segment_count):
        end = duration if index == segment_count - 1 else round(step * (index + 1), 1)
        ranges.append((start, end))
        start = end
    return ranges


def movement_keywords(text: str) -> bool:
    return any(k in text for k in ["显化", "生出", "生成", "降下", "落下", "消散", "化成", "吸入", "收束", "变身", "爆发", "产品出场", "出场"])


def frame_call(index: int, count: int) -> str:
    if count == 1:
        return "00:00 使用首帧图开始生成；本镜头为简单短镜头，沿用首帧图连续生成。"
    if index == 0:
        return "本段从首帧图开始生成，首帧图作为 00:00 起始状态。"
    if index == count - 1:
        if count == 2:
            return "本段从首帧图的自然运动过渡至结尾帧图，结尾帧图锁定镜头结束状态。"
        return "本段从中间过程自然过渡至结尾帧图，结尾帧图锁定镜头结束状态。"
    return "本段参考中间关键帧图，锁定动作过程、主体位置和构图变化。"


def transition_note(index: int, count: int) -> str:
    if index == count - 1:
        return "动作在结尾帧状态稳定收束；下一镜头从该结果自然承接，如跨场景则在此处直接切镜。"
    return "保持人物、产品、道具位置和动作方向连续，自然进入下一时间段。"


def inline_asset_tags(refs: str) -> str:
    tags: list[str] = []
    seen: set[str] = set()
    for raw_line in refs.splitlines():
        line = raw_line.strip().lstrip("-").strip()
        if not line or "：" not in line:
            continue
        label, filename = [part.strip() for part in line.split("：", 1)]
        if not filename or filename == "待补充":
            continue
        if label == "宴会厅" or "宴会厅全景" in filename:
            tag = f"背景@{filename}"
        elif filename.startswith("角色_"):
            tag = f"人物@{filename}"
        elif filename.startswith("场景_"):
            tag = f"场景@{filename}"
        elif filename.startswith("道具_"):
            tag = f"道具@{filename}"
        elif filename.startswith("特效_"):
            tag = f"特效@{filename}"
        elif "产品" in label or "空调" in label or "柜机" in filename:
            tag = f"产品@{filename}"
        elif "群像" in label or "群像" in filename:
            tag = f"人物群像@{filename}"
        else:
            tag = f"素材@{filename}"
        if tag not in seen:
            seen.add(tag)
            tags.append(tag)
    return "；".join(tags) if tags else "沿用本镜头已确认设定"


def timecoded_content(row: dict[str, str], refs: str) -> str:
    text = row.get("picture", "")
    duration = parse_duration_seconds(row.get("duration", ""))
    transformation = movement_keywords(text)
    ranges = split_ranges(duration, 3 if transformation else 2)
    asset_anchor = inline_asset_tags(refs)
    subject = subject_hint(text)

    if transformation:
        if any(k in text for k in ["消散", "化成", "吸入", "收束"]):
            stage_names = ["转化开始", "路径收束", "结果定格"]
            action_notes = [
                "主体从原始状态开始变化，先让黑烟/黑风/气流的方向和起点清楚。",
                "变化沿明确轨迹推进，烟雾或气流被引导到目标位置，过程不要遮挡关键产品结构。",
                "变化完成，结果稳定，产品或关键主体保持清晰可识别。",
            ]
            effect_notes = [
                "第一层烟雾、风流或能量线出现，形态真实，不做恐怖腐烂表现。",
                "粒子、烟雾、水流或光线形成可读的运动路径，注意前后景层次。",
                "特效逐渐减弱，画面干净，给剪辑留出结束点。",
            ]
        else:
            stage_names = ["显化起点", "生成过程", "完成露出"]
            action_notes = [
                "主体和动作起点明确，角色姿态、手部位置、产品/道具尚未完全出现。",
                "产品或法术效果逐步出现，运动方向、速度和主体关系清楚。",
                "产品或关键主体完整露出并稳定，形成可用的广告画面。",
            ]
            effect_notes = [
                "光源、能量、水汽或风流先作为前兆出现，不要一次性完成变化。",
                "特效围绕主体生成，材质真实，避免遮挡产品结构和Logo。",
                "特效收干净，最终形态清楚，画面不过曝不过暗。",
            ]
        blocks: list[str] = []
        for index, (start, end) in enumerate(ranges):
            stage = stage_names[min(index, len(stage_names) - 1)]
            action = action_notes[min(index, len(action_notes) - 1)]
            effect = effect_notes[min(index, len(effect_notes) - 1)]
            blocks.append(
                f"{format_time(start)}-{format_time(end)} · {stage}\n"
                f"景别：{shot_style(row)[0].split('景别：', 1)[1].split('。', 1)[0]}。\n"
                f"机位：平视机位；{subject}显化、吸入或动作强化阶段可轻微低角度，脸部/产品结构不出画。\n"
                f"构图：以{subject}和变化路径为画面主轴，风、光、烟或人物运动的方向在竖屏内清楚可读；画面锚点：{asset_anchor}。\n"
                f"运镜手法：镜头从{subject}的动作起点克制推近或短距离跟随，速度慢于主体变化，避免无目的晃动。\n"
                f"动作：{action}结合执行内容：{text}\n"
                f"特效：{effect}\n"
                "声音：环境声、能量声或产品风声随动作增强，不盖过台词。\n"
                f"结束状态：{subject}完成本阶段变化并稳定停留在可衔接下一阶段的位置。\n"
                f"衔接要求：{transition_note(index, len(ranges))}"
            )
        return "\n\n".join(blocks)

    labels = ["分镜一", "分镜二", "分镜三", "分镜四"]
    blocks = []
    for index, (start, end) in enumerate(ranges):
        if duration and duration <= 2:
            content = text
        elif index == 0:
            content = f"建立主体、场景和动作起点，执行内容：{text}"
        elif index == len(ranges) - 1:
            content = f"动作推进到结果，保留清晰结束状态，执行内容：{text}"
        else:
            content = f"动作连续推进，保持角色、产品和空间关系一致，执行内容：{text}"
        blocks.append(
            f"{labels[min(index, len(labels) - 1)]}：{format_time(start)}-{format_time(end)}\n"
            f"景别：{shot_style(row)[0].split('景别：', 1)[1].split('。', 1)[0]}。\n"
            f"机位：平视机位；{subject}情绪强化时可略低角度，眼神方向和手部动作必须清楚。\n"
            f"构图：{subject}占据竖屏主要阅读区，动作方向一侧保留可见余量；场景信息通过屏风、桌面、光源或产品位置交代。画面锚点：{asset_anchor}。\n"
            f"运镜手法：镜头跟随{subject}的表情或动作节奏做轻微推近/停顿，不做无目的摇晃。\n"
            f"动作：{content}\n"
            "特效：无；如原脚本要求环境风、光线或粒子，则仅保留与动作直接相关的克制效果。\n"
            "声音：只生成对白、同期环境声和动作音效；不生成音乐/BGM，音乐后期单独配。\n"
            f"结束状态：{subject}完成当前阶段动作，停留在清晰可剪辑的结束姿态。\n"
            f"衔接要求：{transition_note(index, len(ranges))}"
        )
    return "\n".join(blocks)


def long_dialogue_or_performance(row: dict[str, str]) -> bool:
    dialogue = row.get("dialogue", "")
    text = row.get("picture", "")
    duration = parse_duration_seconds(row.get("duration", "")) or 0
    return duration >= 5 and (len(dialogue) >= 18 or any(k in text for k in ["台词", "说出", "对白", "表情", "眼神", "欠身", "执扇", "口播"]))


def keyframe_plan(row: dict[str, str]) -> dict[str, bool | int]:
    text = row.get("picture", "")
    duration = parse_duration_seconds(row.get("duration", ""))
    movement = movement_keywords(text)
    middle_count = 0
    if movement and (duration is None or duration > 2):
        middle_count = 1
    if long_dialogue_or_performance(row):
        middle_count = 2 if (duration or 0) >= 6 else 1
    if duration and duration > 7 and middle_count < 2:
        middle_count = 2
    end_needed = bool((duration and duration > 2) or movement)
    return {
        "first": True,
        "middle": middle_count > 0,
        "middle_count": middle_count,
        "end": end_needed,
    }


def keyframe_decision_text(row: dict[str, str]) -> str:
    plan = keyframe_plan(row)
    duration = parse_duration_seconds(row.get("duration", ""))
    text = row.get("picture", "")
    reasons: list[str] = []
    if duration is not None:
        reasons.append(f"时长 {duration:g}s")
    if movement_keywords(text):
        reasons.append("含生成/变化/转化类动作")
    elif duration and duration > 4:
        reasons.append("长台词或多节奏动作，需要控制过程状态")
    elif duration and duration <= 2:
        reasons.append("短简单镜头，优先首帧连续生成")
    else:
        reasons.append("普通动作推进，按首尾状态控制")
    need = ["首帧"]
    if plan["middle"]:
        middle_count = int(plan.get("middle_count", 1))
        need.append(f"中间关键帧{middle_count}张" if middle_count > 1 else "中间关键帧")
    if plan["end"]:
        need.append("结尾帧")
    skip = []
    if not plan["middle"]:
        skip.append("中间关键帧")
    if not plan["end"]:
        skip.append("结尾帧")
    lines = [
        "视频Prompt先行：先按原始/确认脚本写逐秒视频内容，再从动作结构、结束状态和生成前执行说明反推本镜头需要哪些图片。",
        f"本镜头生成：{'、'.join(need)}。",
    ]
    if skip:
        lines.append(f"本镜头跳过：{'、'.join(skip)}。")
    lines.append(f"判定依据：{'；'.join(reasons)}。")
    lines.append("相邻复用：未判定为与上一镜头结尾帧完全相同；如后续人工确认连续画面完全一致，可将本镜头首帧改为复用上一镜头结尾帧。")
    return "\n".join(lines)


def shot_number_text(value: str) -> str:
    try:
        return f"{int(float(str(value))):02d}"
    except ValueError:
        return str(value).strip().zfill(2)


def execution_note_for(row: dict[str, str], refs: str) -> str:
    raise RuntimeError("脚本不允许生成 `生成前执行说明`。该列必须由AI通读全片后逐镜头填写。")


def keyframe_call_section_for(row: dict[str, str]) -> str:
    shot = shot_number_text(row.get("shot", ""))
    plan = keyframe_plan(row)
    first_file = f"镜头{shot}_首帧_v01.png"
    middle_count = int(plan.get("middle_count", 0))
    if middle_count <= 0:
        middle_lines = ["中间关键帧：不使用单独中间关键帧，以分镜动作、表情和运镜控制中段。"]
    elif middle_count == 1:
        middle_lines = [f"中间关键帧：镜头{shot}_中间关键帧_v01.png，锁定本镜头中段的主体位置、动作过程和构图变化。"]
    else:
        middle_lines = [
            f"中间关键帧{chr(65 + index)}：镜头{shot}_中间关键帧{chr(65 + index)}_v01.png，控制第{index + 1}个中段表演/动作状态。"
            for index in range(middle_count)
        ]
    if plan["end"]:
        end_line = f"结尾帧：镜头{shot}_结尾帧_v01.png，锁定镜头结束状态；末段动作收束到该画面。"
    else:
        end_line = "结尾帧：不使用单独结尾帧，结束状态以【画面内容】最后一个时间段控制。"
    return "\n".join([
        "【关键帧调用】",
        f"首帧：{first_file}，作为00:00起始状态。",
        *middle_lines,
        end_line,
    ])


def disabled_keyframe_prompt(frame_name: str, reason: str) -> str:
    return ""


def frame_prompts(row: dict[str, str], continuity: str) -> tuple[str, str, str]:
    raise RuntimeError("脚本不允许生成首帧/中间关键帧/结尾帧 Prompt。Prompt 必须由AI逐镜头撰写。")


def build_prompt(row: dict[str, str], refs: str, first: str, middle: str, end: str, continuity: str, project_style: str) -> str:
    raise RuntimeError("脚本不允许生成视频内容 Prompt。Prompt 必须由AI通读全片后逐镜头撰写。")


def estimate_row_height(row, header_row: int) -> float:
    if row[0].row == header_row:
        return 34
    max_chars = 0
    max_lines = 1
    for cell in row:
        text = str(cell.value or "")
        max_chars = max(max_chars, len(text))
        max_lines = max(max_lines, text.count("\n") + 1)
    # Keep rows scannable. Long Prompt cells remain wrapped and inspectable in the formula bar.
    by_length = 54 + min(max_chars / 95, 6) * 16
    by_lines = min(max_lines, 8) * 15
    return max(68, min(168, by_length, max(72, by_lines)))


def style_project_info_sheet(ws, title: str = "项目说明"):
    title_fill = PatternFill("solid", fgColor="1F4E78")
    label_fill = PatternFill("solid", fgColor="EAF3F8")
    thin = Side(style="thin", color="D9E2F3")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.sheet_view.zoomScale = 90
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 96
    ws.freeze_panes = "A2"

    ws.insert_rows(1)
    ws["A1"] = title
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=2)
    ws["A1"].fill = title_fill
    ws["A1"].font = Font(name="Microsoft YaHei", size=13, bold=True, color="FFFFFF")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    for row in ws.iter_rows(min_row=2, max_col=2):
        for cell in row:
            cell.border = border
            cell.font = Font(name="Microsoft YaHei", size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        row[0].fill = label_fill
        row[0].font = Font(name="Microsoft YaHei", size=10, bold=True, color="1F4E78")
        row[0].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        row[1].alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        text = str(row[1].value or "")
        row_idx = row[0].row
        ws.row_dimensions[row_idx].height = max(28, min(96, 24 + text.count("\n") * 16 + len(text) / 90 * 14))


def style_sheet(ws, widths: list[int], header_row: int = 1):
    header_fill = PatternFill("solid", fgColor="1F4E78")
    subtle_fill = PatternFill("solid", fgColor="D9EAF7")
    thin = Side(style="thin", color="D9E2F3")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for cell in ws[header_row]:
        cell.fill = header_fill
        cell.font = Font(name="Microsoft YaHei", color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border

    for row in ws.iter_rows(min_row=header_row + 1):
        for cell in row:
            cell.font = Font(name="Microsoft YaHei", size=10)
            cell.alignment = body_alignment(ws.cell(header_row, cell.column).value)
            cell.border = border
        if row[0].value == "合计":
            for cell in row:
                cell.fill = subtle_fill
                cell.font = Font(name="Microsoft YaHei", size=10, bold=True)

    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = width
    for row_index in range(1, ws.max_row + 1):
        ws.row_dimensions[row_index].height = estimate_row_height(ws[row_index], header_row)
    ws.freeze_panes = f"A{header_row + 1}"
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.zoomScale = 80


def default_output_path(input_path: Path, project_dir: Path | None) -> Path:
    filename = f"{input_path.stem}_内部执行脚本与Prompt表_v01.xlsx"
    if project_dir:
        target_dir = project_dir / "05_内部制作执行"
        if not target_dir.exists():
            target_dir = project_dir / "06_内部制作执行"
        if not target_dir.exists():
            target_dir = project_dir / "04_分镜与视频生成"
        return target_dir / filename
    return input_path.with_name(filename)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True, help="Confirmed execution script workbook")
    parser.add_argument("--output", help="Output prompt workbook path")
    parser.add_argument("--project-dir", help="Optional project root folder")
    args = parser.parse_args()

    input_path = Path(args.input).expanduser()
    project_dir = Path(args.project_dir).expanduser() if args.project_dir else find_project_root(input_path)
    output_path = Path(args.output).expanduser() if args.output else default_output_path(input_path, project_dir)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    rows = read_execution_rows(input_path)
    assets = scan_assets(project_dir)
    continuity = infer_continuity(rows)
    project_style = infer_project_style(rows)

    wb = Workbook()
    ws = wb.active
    ws.title = "执行脚本与Prompt表"
    info_ws = wb.create_sheet("项目说明")
    meta = [
        ["项目", project_dir.name.replace("_项目交付文件夹", "") if project_dir else input_path.stem],
        ["文件用途", "内部制作执行表结构：同一行查看原始脚本、确认后执行脚本和参考素材；Prompt列留空，等待AI导演逐镜头撰写。"],
        ["画面规格", "9:16竖屏"],
        ["Prompt标准", "脚本不生成Prompt。AI通读全片后先锁定连续性，再逐镜头填写首帧Prompt、中间关键帧Prompt、结尾帧Prompt、生成前执行说明和视频内容Prompt。"],
        ["执行口径", "分镜图不再提交用户确认；本脚本仅生成结构和参考素材，不进入首帧图或视频生成。"],
        ["版本", "v01"],
    ]
    for meta_row in meta:
        info_ws.append(meta_row)
    ws.append(["镜号", "原始镜号", "原始脚本内容", "确认后执行内容", "执行时长(s)", "台词/旁白", "音效/音乐", "参考素材", "首帧Prompt", "中间关键帧Prompt", "结尾帧Prompt", "生成前执行说明", "视频内容Prompt（含声音/负面）", "用户确认/修改依据", "样片建议", "制作状态", "输出文件名", "备注"])

    for row in rows:
        if row.get("total"):
            ws.append(["合计", "", "", "", row.get("duration"), "", "", "", "", "", "", "", "", "", "", "", "", ""])
            continue
        refs = references_for(row, assets)
        ws.append(
            [
                row.get("shot"),
                row.get("orig"),
                row.get("orig_text"),
                row.get("picture"),
                row.get("duration"),
                row.get("dialogue"),
                row.get("sfx"),
                refs,
                "",
                "",
                "",
                "",
                "",
                row.get("note"),
                "",
                "待AI逐镜头撰写Prompt",
                "",
                "脚本仅生成结构；Prompt列必须由AI通读全片后填写。",
            ]
        )

    style_sheet(ws, [7, 8, 36, 32, 10, 15, 16, 36, 50, 34, 34, 50, 78, 24, 18, 11, 18, 18], header_row=1)
    style_project_info_sheet(info_ws)
    wb.save(output_path)
    print(output_path)


if __name__ == "__main__":
    main()
