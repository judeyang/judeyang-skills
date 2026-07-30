#!/usr/bin/env python3
"""Legacy ReportLab fallback for a client-facing final confirmation PDF.

Normal final confirmation and closeout PDFs must be generated through the
tujinpdf HTML + browser-rendering workflow. This script is retained only for
explicit legacy fallback use.
"""

from __future__ import annotations

import argparse
import subprocess
import sys
import tempfile
from pathlib import Path

def latest_file(folder: Path, pattern: str) -> Path:
    files = sorted(folder.glob(pattern))
    if not files:
        raise FileNotFoundError(f"未找到文件：{folder / pattern}")
    return files[-1]


def safe_text(value: object) -> str:
    return "" if value is None else str(value).strip()


def collect_confirmation_rows(workbook: Path) -> list[tuple[str, str, str]]:
    ws = load_workbook(workbook, data_only=True)["脚本审核确认表"]
    headers = {safe_text(cell.value): cell.column for cell in ws[1]}
    required = ["对应镜号/范围", "问题类型", "用户意见/选择"]
    for name in required:
        if name not in headers:
            raise RuntimeError(f"确认表缺少列：{name}")

    rows: list[tuple[str, str, str]] = []
    for row in range(2, ws.max_row + 1):
        scope = safe_text(ws.cell(row, headers["对应镜号/范围"]).value)
        issue = safe_text(ws.cell(row, headers["问题类型"]).value)
        decision = safe_text(ws.cell(row, headers["用户意见/选择"]).value)
        if not scope or issue == "无意见":
            continue
        decision = decision.replace("默认无意见，按当前方案执行", "未提出调整的内容按当前方案执行")
        rows.append((scope, issue, decision or "未提出调整的内容按当前方案执行"))
    return rows[:18]


def collect_asset_images(project_dir: Path) -> list[Path]:
    root = project_dir / "03_设定资产"
    if not root.exists():
        root = project_dir / "04_设定资产"
    images = sorted(
        path
        for path in root.rglob("*")
        if path.is_file() and path.suffix.lower() in {".png", ".jpg", ".jpeg", ".webp"}
        and not path.name.startswith(".")
        and "旧参考" not in path.stem
    )
    seat_images = [path for path in images if "贾母C位座次" in path.stem]
    if not seat_images:
        return images

    preferred_names = [
        "场景_贾母C位座次_v13_Codex人物参考重绘_9比16.png",
        "场景_贾母C位座次_v13_Codex人物参考重绘.png",
        "场景_贾母C位座次_v13_人物参考重绘.png",
        "场景_贾母C位座次_预览_v02.png",
    ]
    selected_seat = None
    by_name = {path.name: path for path in seat_images}
    for name in preferred_names:
        if name in by_name:
            selected_seat = by_name[name]
            break
    selected_seat = selected_seat or seat_images[-1]
    return [path for path in images if "贾母C位座次" not in path.stem] + [selected_seat]


def display_asset_name(path: Path) -> str:
    name = path.stem
    if "贾母C位座次" in name:
        return "场景_贾母C位座次_确认版"
    replacements = [
        "_设定_预览_v02",
        "_关键帧_预览_v02",
        "_氛围设定_预览_v02",
        "_预览_v02",
        "_确认版",
    ]
    for item in replacements:
        name = name.replace(item, "")
    return name


def product_materials_received(project_dir: Path) -> bool:
    folder = project_dir / "01_客户原始资料" / "02_产品素材"
    return folder.exists() and any(
        path.is_file() and not path.name.startswith(".")
        for path in folder.rglob("*")
    )


def optimized_gallery_image(source: Path, cache_dir: Path) -> Path:
    target = cache_dir / f"{source.stem}.jpg"
    with PILImage.open(source) as image:
        image = image.convert("RGB")
        image.thumbnail((1800, 1800), PILImage.Resampling.LANCZOS)
        image.save(target, "JPEG", quality=84, optimize=True)
    return target


def paragraph(text: str, style: ParagraphStyle) -> Paragraph:
    return Paragraph(text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"), style)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--project-dir", required=True)
    parser.add_argument("--project-name", required=True)
    parser.add_argument("--output", required=True)
    parser.add_argument("--confirm-date", default="")
    parser.add_argument("--project-stage", default="进入样片及后续视频制作阶段")
    parser.add_argument("--sample-status", default="")
    parser.add_argument(
        "--allow-legacy-reportlab",
        action="store_true",
        help="Required. Use only when tujinpdf is unavailable and the user explicitly accepts the legacy ReportLab fallback.",
    )
    args = parser.parse_args()

    if not args.allow_legacy_reportlab:
        print(
            "Refusing to generate a formal client-facing PDF with the legacy ReportLab fallback. "
            "Use the tujinpdf HTML + browser-rendering workflow, or rerun with "
            "--allow-legacy-reportlab only after explicit user approval.",
            file=sys.stderr,
        )
        return 2

    global load_workbook, PILImage, colors, TA_CENTER, TA_LEFT, A4
    global ParagraphStyle, getSampleStyleSheet, mm, pdfmetrics, UnicodeCIDFont
    global Image, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    from openpyxl import load_workbook
    from PIL import Image as PILImage
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.platypus import (
        Image,
        PageBreak,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    project_dir = Path(args.project_dir).expanduser()
    output = Path(args.output).expanduser()
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook = latest_file(project_dir / "02_用户确认文件", "*用户已确认*.xlsx")
    audit_rows = collect_confirmation_rows(workbook)
    asset_images = collect_asset_images(project_dir)
    has_product_materials = product_materials_received(project_dir)

    confirm_date = args.confirm_date or "以双方已确认记录为准"
    missing_materials = "如有品牌字体、字幕规范、音乐授权或补充品牌规范，请继续提供。"
    product_status = "已收到产品官方图片及视频素材。产品外观、Logo、屏显与出风口等品牌资产均以贵方提供的官方素材为准。"
    if not has_product_materials:
        product_status = "产品官方图片、视频、Logo、屏显与出风口细节素材待提供。收到后将按官方素材校正产品表现。"

    visible_text: list[str] = []

    def add(text: str) -> str:
        visible_text.append(text)
        return text

    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    styles = getSampleStyleSheet()
    title = ParagraphStyle("TitleCN", parent=styles["Title"], fontName="STSong-Light", fontSize=22, leading=30, alignment=TA_CENTER, textColor=colors.HexColor("#172A46"))
    heading = ParagraphStyle("HeadingCN", parent=styles["Heading1"], fontName="STSong-Light", fontSize=16, leading=22, textColor=colors.HexColor("#172A46"))
    body = ParagraphStyle("BodyCN", parent=styles["BodyText"], fontName="STSong-Light", fontSize=10.5, leading=17, alignment=TA_LEFT, textColor=colors.HexColor("#222222"))
    small = ParagraphStyle("SmallCN", parent=body, fontSize=9, leading=14, textColor=colors.HexColor("#555555"))
    table_body = ParagraphStyle("TableCN", parent=body, fontSize=9, leading=13)
    table_head = ParagraphStyle("TableHeadCN", parent=table_body, alignment=TA_CENTER, textColor=colors.white)

    doc = SimpleDocTemplate(str(output), pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=16 * mm, bottomMargin=16 * mm)
    story = []

    story.extend([
        Spacer(1, 35 * mm),
        paragraph(add(args.project_name), title),
        Spacer(1, 6 * mm),
        paragraph(add("脚本与设定资产最终确认归档"), title),
        Spacer(1, 22 * mm),
        paragraph(add(f"确认时间：{confirm_date}"), body),
        paragraph(add(f"项目阶段：{args.project_stage}"), body),
        PageBreak(),
    ])

    story.extend([
        paragraph(add("确认结论"), heading),
        Spacer(1, 3 * mm),
        paragraph(add("根据已确认内容，项目当前执行方向如下。"), body),
        Spacer(1, 4 * mm),
    ])
    summary_rows = [
        ("脚本审核意见", "已确认。有明确填写的意见按确认内容执行；未提出调整的内容按当前方案执行。"),
        ("形象场景文字方向", "已确认。形象场景方向按当前方案继续。"),
        ("人物、场景、道具与特效设定图", "已单独查看设定图，并确认资产与场景无问题。"),
        ("产品官方素材", product_status),
    ]
    if args.sample_status:
        summary_rows.insert(3, ("样片审核", args.sample_status))
    table = Table([[paragraph(add(k), table_body), paragraph(add(v), table_body)] for k, v in summary_rows], colWidths=[42 * mm, 130 * mm])
    table.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#C8D1DC")),
        ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#EFF3F7")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]))
    story.extend([table, Spacer(1, 5 * mm)])
    story.append(paragraph(add(missing_materials), body))
    story.append(PageBreak())

    story.extend([
        paragraph(add("脚本确认摘要"), heading),
        Spacer(1, 3 * mm),
        paragraph(add("以下为本项目已确认的主要脚本执行方向。"), body),
        Spacer(1, 4 * mm),
    ])
    rows = [[paragraph(add("对应镜号/范围"), table_head), paragraph(add("确认事项"), table_head), paragraph(add("确认结果"), table_head)]]
    rows.extend([[paragraph(add(scope), table_body), paragraph(add(issue), table_body), paragraph(add(decision), table_body)] for scope, issue, decision in audit_rows])
    table = Table(rows, repeatRows=1, colWidths=[28 * mm, 50 * mm, 94 * mm])
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#315A82")),
        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#C8D1DC")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    story.extend([table, Spacer(1, 4 * mm), paragraph(add("未提出调整的内容按当前方案执行。"), body), PageBreak()])

    gallery_cache = tempfile.TemporaryDirectory(prefix="ai-short-video-client-pdf-")
    gallery_cache_dir = Path(gallery_cache.name)
    for index, image_path in enumerate(asset_images, 1):
        story.extend([
            paragraph(add("设定图确认"), heading),
            Spacer(1, 3 * mm),
            paragraph(add(f"{index}/{len(asset_images)}  {display_asset_name(image_path)}"), small),
            Spacer(1, 3 * mm),
        ])
        img = Image(str(optimized_gallery_image(image_path, gallery_cache_dir)))
        max_width, max_height = 170 * mm, 205 * mm
        ratio = min(max_width / img.imageWidth, max_height / img.imageHeight)
        img.drawWidth = img.imageWidth * ratio
        img.drawHeight = img.imageHeight * ratio
        story.extend([img, PageBreak()])

    story.extend([
        paragraph(add("后续执行说明"), heading),
        Spacer(1, 3 * mm),
        paragraph(add("本页用于说明后续制作依据，便于双方沟通确认。"), body),
        Spacer(1, 4 * mm),
        paragraph(add("1. 后续视频制作以已确认脚本、形象场景方向和本归档中的设定图为参考依据。") if args.sample_status else add("1. 后续样片与视频制作以已确认脚本、形象场景方向和本归档中的设定图为参考依据。"), body),
        paragraph(add("2. 如后续需调整已确认内容，将同步记录对应调整范围，便于双方沟通确认。"), body),
        paragraph(add("3. 产品画面将按已收到的官方素材校正产品外观、Logo、屏显和出风口等细节。") if has_product_materials else add("3. 产品官方素材提供后，将按官方素材校正产品外观、Logo、屏显和出风口等细节。"), body),
        paragraph(add("4. 样片审核已通过，后续视频制作将按已确认方向推进。") if args.sample_status else add("4. 后续视频制作将按已确认方向推进。"), body),
        Spacer(1, 12 * mm),
        paragraph(add("确认说明：本归档所列脚本方向、形象场景方向及设定图作为后续制作参考依据。"), body),
    ])

    manifest = output.with_suffix(".manifest.txt")
    manifest.write_text("\n".join(visible_text) + "\n", encoding="utf-8")
    validator = Path(__file__).with_name("validate_client_facing_text.py")
    subprocess.run([sys.executable, str(validator), "--input", str(manifest)], check=True)
    doc.build(story)
    gallery_cache.cleanup()
    print(f"PASS: 已生成客户确认 PDF：{output}")
    print(f"PASS: 已生成客户文案清单：{manifest}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
