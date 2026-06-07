#!/usr/bin/env python3
"""Generate a client-facing script audit workbook from a script Excel file."""

from __future__ import annotations

import argparse
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


SHOT_SYNONYMS = ["镜号", "镜头号", "分镜", "序号", "shot"]
VIEW_SYNONYMS = ["景别", "镜头景别", "画幅"]
PICTURE_SYNONYMS = ["画面内容", "画面", "内容", "视频画面", "镜头内容"]
DURATION_SYNONYMS = ["时长", "时间", "秒数", "duration"]
DIALOGUE_SYNONYMS = ["台词/旁白", "台词", "旁白", "文案", "字幕", "口播"]
SFX_SYNONYMS = ["音效/音乐", "音效", "音乐", "声音"]
NOTE_SYNONYMS = ["备注", "说明", "执行备注", "备注说明"]

CENTER_HEADERS = {
    "序号",
    "镜号",
    "对应镜号/范围",
    "原时长",
    "时长",
    "景别",
    "类别",
    "项目",
    "确认模块",
    "问题类型",
    "风险等级",
    "严重度",
    "用户确认",
    "制作状态",
}


def body_alignment(header_value):
    header = str(header_value or "").strip()
    if header in CENTER_HEADERS or header.endswith("编号"):
        return Alignment(horizontal="center", vertical="top", wrap_text=True)
    return Alignment(horizontal="left", vertical="top", wrap_text=True)


@dataclass
class Shot:
    number: str
    view: str = ""
    picture: str = ""
    duration: str = ""
    dialogue: str = ""
    sfx: str = ""
    note: str = ""


def normalize(value) -> str:
    return str(value or "").strip().replace("\n", " / ")


def find_header_row(ws) -> tuple[int, dict[str, int]] | None:
    for row_idx in range(1, min(ws.max_row, 30) + 1):
        values = [normalize(ws.cell(row_idx, col).value).lower() for col in range(1, ws.max_column + 1)]
        if not any(values):
            continue

        mapping: dict[str, int] = {}
        for col_idx, value in enumerate(values, 1):
            raw = value.replace(" ", "")
            if any(k.lower() in raw for k in SHOT_SYNONYMS) and "shot" not in mapping:
                mapping["shot"] = col_idx
            if any(k.lower() in raw for k in VIEW_SYNONYMS) and "view" not in mapping:
                mapping["view"] = col_idx
            if any(k.lower() in raw for k in PICTURE_SYNONYMS) and "picture" not in mapping:
                mapping["picture"] = col_idx
            if any(k.lower() in raw for k in DURATION_SYNONYMS) and "duration" not in mapping:
                mapping["duration"] = col_idx
            if any(k.lower() in raw for k in DIALOGUE_SYNONYMS) and "dialogue" not in mapping:
                mapping["dialogue"] = col_idx
            if any(k.lower() in raw for k in SFX_SYNONYMS) and "sfx" not in mapping:
                mapping["sfx"] = col_idx
            if any(k.lower() in raw for k in NOTE_SYNONYMS) and "note" not in mapping:
                mapping["note"] = col_idx

        if "shot" in mapping and ("picture" in mapping or "dialogue" in mapping):
            return row_idx, mapping

    return None


def parse_duration_seconds(value: str) -> float | None:
    match = re.search(r"(\d+(?:\.\d+)?)", value or "")
    return float(match.group(1)) if match else None


def effective_dialogue_chars(text: str) -> int:
    text = text or ""
    if "标题字" in text and not re.search(r"[：:]\s*[^/]+(?:说|道|曰)", text):
        return 0
    text = re.sub(r"（[^）]*）", "", text)
    text = re.sub(r"\([^)]*\)", "", text)
    text = re.sub(r"【[^】]*】", "", text)
    text = re.sub(r"[^0-9A-Za-z\u4e00-\u9fff]", "", text)
    text = re.sub(r"^(孙悟空|悟空|八戒|猪八戒|唐三藏|唐僧|沙僧|黄风怪|旁白|画外音)", "", text)
    return len(text)


def read_shots(path: Path) -> list[Shot]:
    wb = load_workbook(path, data_only=False)
    shots: list[Shot] = []

    for ws in wb.worksheets:
        found = find_header_row(ws)
        if found:
            header_row, mapping = found
            for row_idx in range(header_row + 1, ws.max_row + 1):
                shot_no = normalize(ws.cell(row_idx, mapping["shot"]).value)
                if not shot_no:
                    continue
                if shot_no in {"合计", "总计"}:
                    continue
                shots.append(
                    Shot(
                        number=shot_no,
                        view=normalize(ws.cell(row_idx, mapping.get("view", 0)).value) if mapping.get("view") else "",
                        picture=normalize(ws.cell(row_idx, mapping.get("picture", 0)).value) if mapping.get("picture") else "",
                        duration=normalize(ws.cell(row_idx, mapping.get("duration", 0)).value) if mapping.get("duration") else "",
                        dialogue=normalize(ws.cell(row_idx, mapping.get("dialogue", 0)).value) if mapping.get("dialogue") else "",
                        sfx=normalize(ws.cell(row_idx, mapping.get("sfx", 0)).value) if mapping.get("sfx") else "",
                        note=normalize(ws.cell(row_idx, mapping.get("note", 0)).value) if mapping.get("note") else "",
                    )
                )
        else:
            # Fallback for common shot-table layout:
            # A=shot, B=view, C=picture, D=duration, E=dialogue, F=sfx, G=note
            for row_idx in range(1, ws.max_row + 1):
                first = normalize(ws.cell(row_idx, 1).value)
                if not first or not re.match(r"^\d+[A-Za-z]?$", first):
                    continue
                shots.append(
                    Shot(
                        number=first,
                        view=normalize(ws.cell(row_idx, 2).value),
                        picture=normalize(ws.cell(row_idx, 3).value),
                        duration=normalize(ws.cell(row_idx, 4).value),
                        dialogue=normalize(ws.cell(row_idx, 5).value),
                        sfx=normalize(ws.cell(row_idx, 6).value),
                        note=normalize(ws.cell(row_idx, 7).value),
                    )
                )

    # Preserve order but remove exact duplicate shot rows.
    seen = set()
    unique: list[Shot] = []
    for shot in shots:
        key = (shot.number, shot.picture, shot.dialogue)
        if key not in seen:
            unique.append(shot)
            seen.add(key)
    return unique


def original_summary(shot: Shot) -> str:
    parts = []
    if shot.view:
        parts.append(f"景别：{shot.view}")
    if shot.picture:
        parts.append(f"画面：{shot.picture}")
    if shot.dialogue:
        parts.append(f"台词/旁白：{shot.dialogue}")
    if shot.note:
        parts.append(f"备注：{shot.note}")
    return "\n".join(parts)


def base_confirm_rows() -> list[list[str]]:
    return [
        [
            "制作前确认",
            "",
            "客户原脚本是否必须100%忠于执行？",
            "客户权限确认",
            "高",
            "如果客户要求逐字逐镜执行，审核只能提示风险；如果允许专业调整，才能提前修正逻辑、台词时长和AI生成风险。",
            "建议明确哪些必须保留，哪些允许为成片质量微调。",
            "A 100%忠于原脚本；B 保留核心设定，允许专业微调；C 可重构但保留品牌卖点。建议：B。",
            "",
        ],
        [
            "制作前确认",
            "",
            "是否接受增加时长？",
            "建议增加时长",
            "高",
            "如果信息量超过原定时长，硬塞会影响表演、产品露出、特效完成度和成片质感。",
            "建议将增加时长作为正式选项，而不是只在短时长内压缩。",
            "A 保持原时长；B 增加15-30秒；C 删除部分内容保持原时长。建议：B。",
            "",
        ],
        [
            "制作前确认",
            "",
            "最终视频画面比例和横竖屏如何确认？",
            "画面比例确认",
            "高",
            "画面比例会直接影响构图、人物站位、产品露出、字幕安全区和最终导出规格，后期再改容易造成返工。",
            "建议制作前按投放平台确认明确比例；短视频平台通常优先9:16竖屏，横版发布或大屏播放优先16:9横屏。",
            "A 9:16竖屏；B 16:9横屏；C 仅确认横屏，比例待定；D 仅确认竖屏，比例待定。建议按投放平台选择A或B。",
            "",
        ],
        [
            "制作前确认",
            "",
            "哪些内容属于客户不可改的硬要求？",
            "边界确认",
            "中",
            "需要先确认品牌名、产品名、卖点、角色、台词、时长、风格等不可改内容。",
            "请客户列出不可改项，避免专业修改误伤客户重点。",
            "A 客户列出不可改项；B 默认仅保留品牌和卖点；C 全部可讨论。建议：A。",
            "",
        ],
        [
            "制作前确认",
            "",
            "产品、Logo、字体、音乐、配音等素材是否齐全？",
            "媒体资产确认",
            "高",
            "产品外观、Logo、字体授权、落版规范不清，会导致后期返工或品牌错误。",
            "建议制作前收齐官方产品图、Logo源文件、品牌色、卖点标准字、字体授权、音乐/配音要求。",
            "A 甲方提供完整素材；B 暂用占位素材预览；C AI自由生成。建议：A。",
            "",
        ],
    ]


def audit_shot(shot: Shot) -> list[list[str]]:
    rows: list[list[str]] = []
    sec = parse_duration_seconds(shot.duration)
    chars = effective_dialogue_chars(shot.dialogue)
    text = f"{shot.picture} {shot.dialogue} {shot.note}"

    if sec and chars:
        cps = chars / sec
        if cps > 5.5:
            rows.append(
                [
                    shot.number,
                    shot.duration,
                    original_summary(shot),
                    "台词超时",
                    "高",
                    f"该镜头约{chars}个有效字，{shot.duration}内口型和表演很难自然完成。",
                    "压缩台词，或延长镜头，或改为画外音。",
                    "A 缩短台词；B 延长镜头；C 改画外音。建议：A。",
                    "",
                ]
            )
        elif cps > 4.2 or (sec <= 2 and chars > 8):
            rows.append(
                [
                    shot.number,
                    shot.duration,
                    original_summary(shot),
                    "台词偏快",
                    "中",
                    f"该镜头约{chars}个有效字，{shot.duration}内偏赶，表演空间不足。",
                    "建议压缩为更短、更口语化的一句。",
                    "A 压缩台词；B 保留台词但延长镜头；C 改旁白。建议：A。",
                    "",
                ]
            )

    complexity_hits = re.findall(r"转场|显化|消散|生长|扩散|海啸|骷髅|水幕|粒子|变成|同时|涌出|落下|钻入|重绘", text)
    if sec and sec <= 3 and len(complexity_hits) >= 2:
        rows.append(
            [
                shot.number,
                shot.duration,
                original_summary(shot),
                "视觉执行风险",
                "高" if sec <= 2 else "中",
                "该镜头在很短时间内包含多个视觉变化，AI生成容易混乱，观众也不容易看清。",
                "建议每个短镜头只保留一个视觉重点，复杂转场拆成多镜。",
                "A 简化视觉重点；B 拆分镜头；C 延长时长。建议：A或B。",
                "",
            ]
        )

    if re.search(r"空调|产品|Logo|品牌|海尔|卡萨帝|统帅", text) and sec and sec <= 4:
        rows.append(
            [
                shot.number,
                shot.duration,
                original_summary(shot),
                "产品露出确认",
                "中",
                "产品或品牌信息出现在短镜头中，需要保证识别时间、外观准确和Logo稳定。",
                "建议使用官方产品素材，并给产品停稳露出的时间。",
                "A 保证产品停稳露出；B 改到结尾英雄镜头；C 仅做背景露出。建议：A。",
                "",
            ]
        )

    if re.search(r"从天而降|直接落下|突然出现|空降", text) and re.search(r"空调|产品|神器", text):
        rows.append(
            [
                shot.number,
                shot.duration,
                original_summary(shot),
                "产品登场生硬",
                "高",
                "现代产品直接出现容易破坏故事世界观，显得像硬插广告。",
                "建议改成世界观内的合理转化，例如法器/道具显化为产品。",
                "A 直接出现；B 法器显化产品；C 角色主动变出产品。建议：B。",
                "",
            ]
        )

    return rows


def no_issue_row(shot: Shot) -> list[str]:
    return [
        shot.number,
        shot.duration,
        original_summary(shot),
        "无意见",
        "低",
        "经逐镜审核，暂未发现明显执行风险，建议保留原脚本内容。",
        "建议保持现有处理；如贵方有补充要求，可在用户意见中说明。",
        "A 保留原内容；B 补充修改意见。建议：A。",
        "",
    ]


def extract_assets(shots: Iterable[Shot]) -> list[list[str]]:
    text = "\n".join(f"{s.picture}\n{s.dialogue}\n{s.note}" for s in shots)
    candidates = [
        ("人物", ["孙悟空", "悟空", "唐三藏", "唐僧", "猪八戒", "八戒", "沙僧", "黄风怪", "灵吉菩萨", "菩萨"]),
        ("动物/坐骑", ["白龙马", "马"]),
        ("场景", ["黄风岭", "花果山", "雨林", "沙丘", "瀑布", "落版"]),
        ("产品/品牌", ["海尔", "卡萨帝", "统帅", "空调", "Logo", "水洗空气"]),
        ("道具/法器", ["金箍棒", "净风宝瓶", "定风珠", "法器", "神器"]),
        ("特效", ["黄沙", "黑烟", "毒烟", "水幕", "祥光", "绿叶", "花瓣", "骷髅"]),
        ("包装文字", ["标题字", "字幕", "口播", "落版"]),
    ]
    rows = [["类别", "项目", "初步设定方向", "是否需要客户提供参考", "需要确认的问题", "用户确认", "用户补充意见/素材链接"]]
    seen = set()
    for category, names in candidates:
        for name in names:
            if name in text and (category, name) not in seen:
                rows.append(
                    [
                        category,
                        name,
                        "根据脚本语境提炼视觉方向，先确认文字描述，再生成设定图。",
                        "如客户有指定参考图/品牌素材请提供；没有则按确认后的文字方向生成。",
                        "造型、风格、尺度、连续性、是否必须参考官方素材。",
                        "",
                        "",
                    ]
                )
                seen.add((category, name))
    return rows


def style_sheet(ws, widths: list[int], header_row: int = 1, row_height: int = 84):
    header_fill = PatternFill("solid", fgColor="1F4E78")
    thin = Side(style="thin", color="D9E2F3")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    severity_fills = {
        "高": PatternFill("solid", fgColor="F4CCCC"),
        "中": PatternFill("solid", fgColor="FFF2CC"),
        "低": PatternFill("solid", fgColor="D9EAD3"),
    }

    for cell in ws[header_row]:
        cell.fill = header_fill
        cell.font = Font(name="Arial", color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border

    for row in ws.iter_rows(min_row=header_row + 1):
        for cell in row:
            cell.font = Font(name="Arial", size=10)
            cell.alignment = body_alignment(ws.cell(header_row, cell.column).value)
            cell.border = border
        if len(row) >= 5 and row[4].value in severity_fills:
            row[4].fill = severity_fills[row[4].value]

    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = width
    for row_index in range(1, ws.max_row + 1):
        ws.row_dimensions[row_index].height = row_height if row_index > header_row else 30
    ws.freeze_panes = f"A{header_row + 1}"
    ws.auto_filter.ref = ws.dimensions


def append_feedback_template(wb: Workbook):
    ws = wb.create_sheet("用户修改意见记录")
    ws.append(
        [
            "反馈时间",
            "反馈来源",
            "关联文件",
            "关联镜头/资产",
            "客户原话",
            "要求类型",
            "是否发生在已确认之后",
            "是否已有样片或图片受影响",
            "需要修改的内容",
            "不能修改的内容",
            "客户提供的新素材",
            "期望完成时间",
            "我的处理建议/备注",
            "处理状态",
            "下一步动作",
        ]
    )
    ws.append(
        [
            "",
            "微信 / 飞书 / 邮件 / 电话 / 会议 / 文件批注 / 其他",
            "",
            "",
            "",
            "脚本修改 / 人物设定 / 场景设定 / 产品素材 / 道具特效 / 配音音乐 / 字幕包装 / 成片节奏 / 其他",
            "是 / 否 / 不确定",
            "是 / 否 / 不确定",
            "",
            "",
            "",
            "",
            "",
            "待处理",
            "",
        ]
    )
    style_sheet(ws, [18, 18, 24, 20, 44, 20, 20, 24, 36, 30, 28, 18, 34, 16, 30], row_height=70)


def sort_issue_key(row: list[str]) -> tuple[int, int]:
    scope = str(row[0] or "")
    if scope == "制作前确认":
        return (0, 0)
    if scope == "整体":
        return (1, 0)
    try:
        return (2, int(float(scope)))
    except ValueError:
        return (3, 9999)


def dedupe_issue_rows(rows: list[list[str]]) -> list[list[str]]:
    """Deduplicate same shot/scope + issue type, preferring richer context."""
    best: dict[tuple[str, str], list[str]] = {}
    order: list[tuple[str, str]] = []

    for row in sorted(rows, key=sort_issue_key):
        key = (str(row[0] or ""), str(row[3] or ""))
        if key not in best:
            best[key] = row
            order.append(key)
            continue

        old = best[key]
        if len(str(row[2] or "")) > len(str(old[2] or "")):
            best[key] = row

    return [best[key] for key in order]


def default_output_path(input_path: Path, project_dir: Path | None) -> Path:
    filename = f"{input_path.stem}_脚本审核确认表_v01.xlsx"
    if project_dir:
        target_dir = project_dir / "02_用户确认文件"
        if not target_dir.exists():
            target_dir = project_dir / "02_确认文件"
        return target_dir / filename
    return input_path.with_name(filename)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True, help="Client script .xlsx file")
    parser.add_argument("--output", help="Output audit workbook path")
    parser.add_argument("--project-dir", help="Optional project root folder")
    args = parser.parse_args()

    input_path = Path(args.input).expanduser()
    project_dir = Path(args.project_dir).expanduser() if args.project_dir else None
    output_path = Path(args.output).expanduser() if args.output else default_output_path(input_path, project_dir)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    shots = read_shots(input_path)
    wb = Workbook()
    ws = wb.active
    ws.title = "脚本审核确认表"
    ws.append(["序号", "对应镜号/范围", "原时长", "原画面/台词", "问题类型", "严重度", "建议", "修改方向", "用户选择项", "用户意见/选择"])

    issue_rows: list[list[str]] = []
    for row in base_confirm_rows():
        issue_rows.append(row)

    for shot in shots:
        shot_rows = audit_shot(shot)
        if shot_rows:
            issue_rows.extend(shot_rows)
        else:
            issue_rows.append(no_issue_row(shot))

    if len(issue_rows) == len(base_confirm_rows()):
        issue_rows.append(["整体", "", "未发现明显执行风险。", "整体确认", "低", "经整体审核，暂未发现明显执行风险。", "建议继续确认剧情逻辑、品牌表达和素材完整度。", "A 通过；B 补充修改意见。建议：A。", ""])

    for index, row in enumerate(dedupe_issue_rows(issue_rows), 1):
        ws.append([index, row[0], row[1], row[2], row[3], row[4], row[5], row[6], row[7], row[8]])

    style_sheet(ws, [8, 14, 10, 42, 16, 10, 42, 42, 36, 20])

    ws_assets = wb.create_sheet("形象场景描述确认表")
    for row in extract_assets(shots):
        ws_assets.append(row)
    style_sheet(ws_assets, [14, 20, 44, 36, 36, 16, 34], row_height=72)
    append_feedback_template(wb)

    wb.save(output_path)
    print(output_path)


if __name__ == "__main__":
    main()
