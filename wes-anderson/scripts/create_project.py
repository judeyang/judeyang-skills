#!/usr/bin/env python3
"""Scaffold an AI short-video production project folder and confirmation workbook."""

from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


BASE_FOLDERS = [
    "01_客户原始资料/01_脚本",
    "01_客户原始资料/02_产品素材",
    "01_客户原始资料/03_用户反馈参考图",
    "01_客户原始资料/04_品牌素材",
    "01_客户原始资料/05_字体授权",
    "02_用户确认文件",
    "03_设定资产",
    "04_最终确认归档",
    "05_内部制作执行/01_首帧与关键帧",
    "05_内部制作执行/02_视频片段",
    "05_内部制作执行/03_成片",
]


CENTER_HEADERS = {
    "序号",
    "镜号",
    "对应镜号/范围",
    "原时长",
    "时长",
    "景别",
    "类别",
    "项目",
    "确认模块",
    "问题类型",
    "风险等级",
    "严重度",
    "用户确认",
    "制作状态",
}


def body_alignment(header_value):
    header = str(header_value or "").strip()
    if header in CENTER_HEADERS or header.endswith("编号"):
        return Alignment(horizontal="center", vertical="top", wrap_text=True)
    return Alignment(horizontal="left", vertical="top", wrap_text=True)


def style_sheet(ws, widths):
    header_fill = PatternFill("solid", fgColor="1F4E78")
    thin = Side(style="thin", color="D9E2F3")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = Font(name="Arial", color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial", size=10)
            cell.alignment = body_alignment(ws.cell(1, cell.column).value)
            cell.border = border

    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = width
    for row_index in range(1, ws.max_row + 1):
        ws.row_dimensions[row_index].height = 70 if row_index > 1 else 30
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def write_confirmation_workbook(path: Path, project_name: str):
    wb = Workbook()

    ws = wb.active
    ws.title = "脚本审核确认表"
    ws.append(["确认模块", "确认项", "当前方案", "风险/说明", "建议", "用户确认", "用户补充意见"])
    script_rows = [
        [
            "执行边界",
            "是否允许专业微调脚本",
            "保留客户核心设定和品牌卖点，允许为视频节奏、逻辑顺畅度、AI生成稳定性做专业微调。",
            "若逐字逐镜执行，可能出现剧情生硬、台词超时、产品登场突兀。",
            "建议允许专业微调，并明确不可改内容。",
            "",
            "",
        ],
        [
            "执行边界",
            "是否接受增加时长",
            "按客户原始目标时长执行，或根据脚本信息密度调整。",
            "信息密度过高会压缩表演、产品露出和特效完成度。",
            "如质量优先，建议允许增加至更合理时长。",
            "",
            "",
        ],
        [
            "执行边界",
            "确认画面比例与横竖屏",
            "制作前需确认最终视频为9:16竖屏、16:9横屏，或仅先确认横屏/竖屏方向。",
            "画面比例会影响构图、人物站位、产品露出、字幕安全区和最终导出规格，后期再改容易返工。",
            "建议按投放平台选择明确比例；短视频平台通常优先9:16竖屏，横版发布或大屏播放优先16:9横屏。",
            "A 9:16竖屏；B 16:9横屏；C 仅确认横屏，比例待定；D 仅确认竖屏，比例待定。建议按投放平台选择A或B。",
            "",
        ],
        [
            "脚本内容",
            "剧情链路是否顺畅",
            "根据脚本审核后整理执行链路。",
            "危机、角色尝试、产品解决方案之间缺少过渡时，成片会生硬。",
            "建议保留核心设定，补足角色主动动作和产品登场逻辑。",
            "",
            "",
        ],
        [
            "脚本内容",
            "台词和镜头时长是否接受调整",
            "短镜头台词按可说完标准压缩。",
            "2-3秒镜头台词过长会导致口型和节奏失败。",
            "建议优先保证表演和镜头可执行。",
            "",
            "",
        ],
    ]
    for row in script_rows:
        ws.append(row)

    ws2 = wb.create_sheet("形象场景描述确认表")
    ws2.append(["类别", "项目", "初步设定方向", "是否需要客户提供参考", "需要确认的问题", "用户确认", "用户补充意见/素材链接"])
    visual_rows = [
        ["人物", "主角/关键角色", "根据脚本提炼，保持角色统一、符合项目调性。", "如客户有指定参考图请提供；没有则按文字方向生成。", "年龄、气质、服装、表情尺度。", "", ""],
        ["场景", "主要场景", "根据脚本提炼，先确认文字方向，再生成场景图。", "如有场景参考图请提供。", "色调、时代感、真实/幻想程度、空间连续性。", "", ""],
        ["产品/品牌", "产品与Logo", "必须以客户官方素材为准，不自由发挥产品结构。", "需要提供产品图、Logo、卖点标准字、落版规范。", "产品角度、Logo位置、屏幕文字、品牌色。", "", ""],
        ["道具/特效", "关键道具和视觉特效", "根据脚本提炼，先确认文字方向，再生成关键帧。", "如有参考图可提供。", "风格、复杂度、是否适合AI生成。", "", ""],
    ]
    for row in visual_rows:
        ws2.append(row)

    ws3 = wb.create_sheet("用户修改意见记录")
    ws3.append(
        [
            "反馈时间",
            "反馈来源",
            "关联文件",
            "关联镜头/资产",
            "客户原话",
            "要求类型",
            "是否发生在已确认之后",
            "是否已有样片或图片受影响",
            "需要修改的内容",
            "不能修改的内容",
            "客户提供的新素材",
            "期望完成时间",
            "我的处理建议/备注",
            "处理状态",
            "下一步动作",
        ]
    )
    ws3.append(
        [
            "",
            "微信 / 飞书 / 邮件 / 电话 / 会议 / 文件批注 / 其他",
            "",
            "",
            "",
            "脚本修改 / 人物设定 / 场景设定 / 产品素材 / 道具特效 / 配音音乐 / 字幕包装 / 成片节奏 / 其他",
            "是 / 否 / 不确定",
            "是 / 否 / 不确定",
            "",
            "",
            "",
            "",
            "",
            "待处理",
            "",
        ]
    )

    style_sheet(ws, [14, 24, 44, 36, 36, 16, 32])
    style_sheet(ws2, [14, 22, 44, 36, 36, 16, 34])
    style_sheet(ws3, [18, 18, 24, 20, 44, 20, 20, 24, 36, 30, 28, 18, 34, 16, 30])
    wb.save(path)


def write_rules(path: Path, project_name: str):
    content = f"""# {project_name} 项目文件夹与命名规则

## 工作流程
1. 先确认脚本：在 `02_用户确认文件` 的脚本确认表中填写用户意见。
2. 如需要提前制作人物、场景、道具等资产，先生成内部 `设定资产制作表_v01.xlsx`，用于你或制作人员按Prompt出图。
3. 资产制作完成后，把实际图片文件名和确认项更新回 `项目名_脚本与资产确认表_v01.xlsx`，再提交用户确认。
4. 用户确认脚本和资产后，进入最终归档 PDF 与内部视频制作阶段。

## 项目接收记录
新项目开始时，应记录客户原始资料的接收时间、资料名称、资料类型、来源说明、处理动作和当前状态。
后续生成 `项目执行过程与修改确认归档` 或 `项目执行总结归档` 时，需要把原始脚本接收时间和关键修改节点体现出来。

## 一级目录说明
- `01_客户原始资料`：客户提供的脚本、产品素材、品牌素材、参考图、字体授权。
- `02_用户确认文件`：需要用户填写、选择、反馈的 Excel，核心文件是 `项目名_脚本与资产确认表_v01.xlsx`。
- `03_设定资产`：根据脚本内容动态生成的人物、场景、道具、特效等设定资产。
- `04_最终确认归档`：客户最终确认 PDF。
- `05_内部制作执行`：最终执行 Prompt 表、首帧/关键帧图、视频片段、成片。Prompt 表直接放在本目录，不单独建立视频 Prompt 文件夹。

## 设定资产目录规则
`03_设定资产` 下的子目录不固定。应先阅读脚本、提炼资产类别，再按项目内容创建。

示例：
- 神话/短剧项目：`01_人物设定图`、`02_场景设定图`、`03_道具设定图`、`04_特效关键帧`
- 汽车广告：`01_车型外观参考`、`02_场景设定图`、`03_驾驶员形象`、`04_动态特效关键帧`
- 美妆广告：`01_人物模特设定`、`02_产品质感参考`、`03_场景氛围图`、`04_质地特效关键帧`

## 命名规则
统一格式：`类型_项目_内容_v版本号.扩展名`

示例：
- `脚本审核确认表_v01.xlsx`
- `项目名_脚本与资产确认表_v01.xlsx`
- `设定资产制作表_v01.xlsx`
- `角色_孙悟空_大头照_无文字_v01.png`
- `角色_孙悟空_全身照_无文字_v01.png`
- `场景_黄风岭_设定图_v01.png`
- `道具_净风宝瓶_设定图_v01.png`
- `关键帧_产品显化_v01.png`
- `归档_客户最终确认_v01.pdf`
- `内部执行脚本与Prompt表_v01.xlsx`

## 规则
- 需要用户填写、选择、反馈的阶段使用 Excel。
- 已确认、需要归档的阶段使用 PDF。
- 内部制作执行使用 Excel。
- 不单独创建 `执行版脚本.xlsx` 或 `形象场景描述确认表.xlsx`。相关内容应合并进 `项目名_脚本与资产确认表_v01.xlsx` 或最终内部执行表。
- 不创建 `提示词参考` 或 `视频Prompt` 空目录；Prompt 内容以最终执行 Excel 为准。
- 不覆盖旧版本文件，修改时递增版本号。
"""
    path.write_text(content, encoding="utf-8")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--project-name", required=True)
    parser.add_argument("--base-dir", default=str(Path.home() / "Desktop"))
    parser.add_argument(
        "--asset-subfolders",
        default="",
        help="Optional comma-separated subfolders under 03_设定资产, e.g. 人物设定图,场景设定图,道具设定图,特效关键帧",
    )
    args = parser.parse_args()

    root = Path(args.base_dir).expanduser() / f"{args.project_name}_项目交付文件夹"
    for folder in BASE_FOLDERS:
        (root / folder).mkdir(parents=True, exist_ok=True)

    asset_names = [item.strip() for item in args.asset_subfolders.split(",") if item.strip()]
    for index, name in enumerate(asset_names, 1):
        (root / "03_设定资产" / f"{index:02d}_{name}").mkdir(parents=True, exist_ok=True)

    write_rules(root / "00_项目说明_文件夹与命名规则.md", args.project_name)
    write_confirmation_workbook(
        root / "02_用户确认文件" / f"{args.project_name}_脚本与资产确认表_v01.xlsx",
        args.project_name,
    )
    print(root)


if __name__ == "__main__":
    main()
