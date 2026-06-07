#!/usr/bin/env python3
"""Validate mandatory per-segment detail labels in an internal video Prompt workbook."""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from openpyxl import load_workbook


REQUIRED_LABELS = ["景别：", "机位：", "构图：", "运镜手法：", "动作：", "特效：", "声音：", "结束状态：", "衔接要求："]
REQUIRED_SAMPLE_STYLE_SECTIONS = ["核心主题：", "【运镜规则】"]
PLACEHOLDERS = ["按本镜头需要选择", "根据主体动作采用", "按画面执行", "保持可识别", "适当", "根据实际情况"]
GENERIC_TEMPLATE_PHRASES = [
    "围绕本段动作组织画面",
    "关键人物、手部动作、道具和视线方向完整入画",
    "视线或运动方向预留空间",
    "视线和动作方向预留空间",
    "按9:16竖屏组织前景、中景、背景层次",
    "固定机位，仅保留轻微呼吸感",
    "主体位于画面视觉中心",
    "主体置于画面视觉中心",
    "前景、中景、背景层次清楚",
    "固定机位或缓慢跟随主体移动",
    "固定机位或缓慢推进",
    "只生成只保留",
    "锁定本镜头中段的动作、表情或视觉变化",
    "锁定中段动作、表情或视觉变化",
    "锁定第1个中段动作或表情状态",
    "锁定第2个中段动作或表情状态",
]
NO_TEXT_RULES = ["无文字纯图片", "禁止生成字幕", "禁止生成字幕、标题", "水印"]
FULL_PATH_MARKERS = ["/" + "Users/", "\\" + "Users\\", "file://"]
MUSIC_PROMPT_MARKERS = ["音乐/音效：", "古筝", "琵琶", "背景乐", "配乐", "合奏", "轮指", "变奏"]
KEYFRAME_FILE_RE = re.compile(
    r"镜头\d+_(?:首帧|中间关键帧[A-Z]?|结尾帧)(?:_[^\\s；;,，。]+)?_v\\d+(?:_[^\\s；;,，。]+)?\\.(?:png|jpg|jpeg|webp)"
)
SOLID_COLOR_KEYFRAME_FILE_RE = re.compile(r"镜头\d+_[^\s；;,，。]*(?:纯黑|纯白|纯色|黑场|白场)[^\s；;,，。]*\.(?:png|jpg|jpeg|webp)")
SOLID_COLOR_KEYFRAME_PHRASES = ["纯黑关键帧", "纯白关键帧", "纯色关键帧", "黑场关键帧", "白场关键帧", "纯黑画面作为关键帧", "纯白画面作为关键帧"]
PRODUCT_RULE_MARKERS = ["机身比例", "Logo位置", "Logo错误", "出风口", "导风板", "屏显", "官方产品素材", "产品结构变形"]
PRODUCT_REF_MARKERS = ["产品@", "CAP", "卡萨帝空调", "空调", "170°", "170度", "气流", "送风", "导风板", "出风口", "屏显"]


def find_header(ws, name: str) -> tuple[int, int]:
    for row in range(1, min(ws.max_row, 30) + 1):
        for col in range(1, ws.max_column + 1):
            if ws.cell(row, col).value == name:
                return row, col
    raise RuntimeError(f"未找到列：{name}")


def find_any_header(ws, names: list[str]) -> tuple[int, int]:
    errors = []
    for name in names:
        try:
            return find_header(ws, name)
        except RuntimeError as exc:
            errors.append(str(exc))
    raise RuntimeError(" / ".join(errors))


def find_prompt_sheet(wb):
    for ws in wb.worksheets:
        try:
            find_header(ws, "视频内容Prompt（含声音/负面）")
            find_header(ws, "镜号")
            return ws
        except RuntimeError:
            continue
    raise RuntimeError("未找到包含 `镜号` 和 `视频内容Prompt（含声音/负面）` 的执行表。")


def extract_blocks(prompt: str) -> list[str]:
    if "【画面内容】" not in prompt:
        return []
    content = prompt.rsplit("【画面内容】", 1)[1]
    if "【声音/台词】" in content:
        content = content.split("【声音/台词】", 1)[0]
    return [block.strip() for block in re.split(r"(?=分镜\d+：|(?=\d{2}:\d{2}(?:\.\d+)?-\d{2}:\d{2}(?:\.\d+)?\s*·))", content) if block.strip()]


def extract_keyframe_call(prompt: str) -> str:
    if "【关键帧调用】" not in prompt:
        return ""
    section = prompt.split("【关键帧调用】", 1)[1]
    match = re.search(r"\n【[^】]+】", section)
    if match:
        section = section[: match.start()]
    return section.strip()


def keyframe_files(text: str) -> list[str]:
    seen: set[str] = set()
    files: list[str] = []
    for filename in KEYFRAME_FILE_RE.findall(text):
        if filename not in seen:
            seen.add(filename)
            files.append(filename)
    return files


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True)
    args = parser.parse_args()

    path = Path(args.input).expanduser()
    wb = load_workbook(path, data_only=True)
    ws = find_prompt_sheet(wb)
    header_row, prompt_col = find_header(ws, "视频内容Prompt（含声音/负面）")
    shot_col = find_header(ws, "镜号")[1]
    _, refs_col = find_any_header(ws, ["参考资产", "参考素材"])
    exec_note_col = find_header(ws, "生成前执行说明")[1]
    first_col = find_header(ws, "首帧Prompt")[1]
    middle_col = find_header(ws, "中间关键帧Prompt")[1]
    end_col = find_header(ws, "结尾帧Prompt")[1]
    errors: list[str] = []

    for row in range(header_row + 1, ws.max_row + 1):
        shot = ws.cell(row, shot_col).value
        if shot in (None, "", "合计"):
            continue
        prompt = str(ws.cell(row, prompt_col).value or "")
        refs_text = str(ws.cell(row, refs_col).value or "")
        exec_note = str(ws.cell(row, exec_note_col).value or "")
        first_prompt = str(ws.cell(row, first_col).value or "")
        middle_prompt = str(ws.cell(row, middle_col).value or "")
        end_prompt = str(ws.cell(row, end_col).value or "")
        keyframe_prompt_text = "\n".join([first_prompt, middle_prompt, end_prompt])
        for phrase in SOLID_COLOR_KEYFRAME_PHRASES:
            if phrase in keyframe_prompt_text or phrase in exec_note:
                errors.append(f"镜头 {shot}: 纯黑/纯白/纯色占位画面不能作为 AI 关键帧生成，请改为视频或后期转场说明")
        for filename in SOLID_COLOR_KEYFRAME_FILE_RE.findall(exec_note + "\n" + prompt):
            errors.append(f"镜头 {shot}: `{filename}` 是纯色/黑场占位帧，不应作为 AI 关键帧接入")
        if first_prompt.startswith("无需单独生成"):
            errors.append(f"镜头 {shot}: 首帧Prompt 不能标记为无需单独生成")
        first_reuses_previous = first_prompt.startswith("复用上一镜头结尾帧")
        needs_middle = "中间关键帧=生成" in exec_note or "中间关键帧：生成" in exec_note
        needs_end = "结尾帧=生成" in exec_note or "结尾帧：生成" in exec_note
        if first_reuses_previous and "复用上一镜头结尾帧" not in exec_note:
            errors.append(f"镜头 {shot}: 首帧复用上一镜头结尾帧，但生成前执行说明没有同步说明复用关系")
        if not needs_middle and middle_prompt.strip():
            errors.append(f"镜头 {shot}: 生成前执行说明未要求中间关键帧，但中间关键帧Prompt未留空")
        if needs_middle and not middle_prompt.strip():
            errors.append(f"镜头 {shot}: 生成前执行说明要求中间关键帧，但中间关键帧Prompt为空")
        if needs_middle and middle_prompt.startswith("无需单独生成"):
            errors.append(f"镜头 {shot}: 生成前执行说明要求中间关键帧，但中间关键帧Prompt被标记为无需生成")
        if not needs_end and end_prompt.strip():
            errors.append(f"镜头 {shot}: 生成前执行说明未要求结尾帧，但结尾帧Prompt未留空")
        if needs_end and not end_prompt.strip():
            errors.append(f"镜头 {shot}: 生成前执行说明要求结尾帧，但结尾帧Prompt为空")
        if needs_end and end_prompt.startswith("无需单独生成"):
            errors.append(f"镜头 {shot}: 生成前执行说明要求结尾帧，但结尾帧Prompt被标记为无需生成")
        if "【全片连续性/空间关系】" not in prompt:
            errors.append(f"镜头 {shot}: 缺少 `【全片连续性/空间关系】`")
        has_product_ref = any(marker in (prompt + "\n" + refs_text + "\n" + exec_note) for marker in PRODUCT_REF_MARKERS)
        if not has_product_ref:
            for marker in PRODUCT_RULE_MARKERS:
                if marker in prompt:
                    errors.append(f"镜头 {shot}: 非产品镜头包含产品专用规则 `{marker}`，请按本镜头相关性裁剪")
        for section in REQUIRED_SAMPLE_STYLE_SECTIONS:
            if section not in prompt:
                errors.append(f"镜头 {shot}: 缺少样板式 Prompt 结构 `{section}`")
        if "关键帧出图" not in exec_note or "视频生成前上传" not in exec_note:
            errors.append(f"镜头 {shot}: 生成前执行说明缺少关键帧出图或视频生成前上传说明")
        if "关键帧视觉复核" not in exec_note:
            errors.append(f"镜头 {shot}: 生成前执行说明缺少关键帧视觉复核要求")
        if "【关键帧调用】" not in prompt:
            errors.append(f"镜头 {shot}: 视频Prompt缺少 `【关键帧调用】`，无法明确首帧/中间关键帧/结尾帧用法")
        call_section = extract_keyframe_call(prompt)
        exec_files = keyframe_files(exec_note)
        call_files = keyframe_files(call_section)
        for filename in exec_files:
            if filename not in call_files:
                errors.append(f"镜头 {shot}: 生成前执行说明中的关键帧 `{filename}` 未出现在视频Prompt的【关键帧调用】中")
        if not needs_middle and "中间关键帧：不使用单独中间关键帧" not in call_section:
            errors.append(f"镜头 {shot}: 生成前执行说明跳过中间关键帧，但【关键帧调用】未明确中间关键帧不使用")
        if not needs_end and "结尾帧：不使用单独结尾帧" not in call_section:
            errors.append(f"镜头 {shot}: 生成前执行说明跳过结尾帧，但【关键帧调用】未明确结尾帧不使用")
        if "引用资产：" in prompt:
            errors.append(f"镜头 {shot}: 视频Prompt包含散乱引用资产执行说明，应移到生成前执行说明并在画面描述中使用 inline asset tags")
        if "音乐/BGM：不生成，后期单独配" not in prompt:
            errors.append(f"镜头 {shot}: 视频Prompt缺少 `音乐/BGM：不生成，后期单独配`")
        for marker in MUSIC_PROMPT_MARKERS:
            if marker in prompt:
                errors.append(f"镜头 {shot}: 视频Prompt包含音乐生成相关表达 `{marker}`，视频工具只生成音效/同期声")
        for marker in FULL_PATH_MARKERS:
            if marker in prompt or marker in refs_text:
                errors.append(f"镜头 {shot}: 参考素材或 Prompt 中包含完整本地路径 `{marker}`，应改为文件名")
        if not all(rule in prompt for rule in NO_TEXT_RULES):
            errors.append(f"镜头 {shot}: 缺少无文字纯图片/禁止字幕水印规则")
        for placeholder in PLACEHOLDERS:
            if placeholder in prompt:
                errors.append(f"镜头 {shot}: 包含占位表达 `{placeholder}`")
        for phrase in GENERIC_TEMPLATE_PHRASES:
            if phrase in prompt:
                errors.append(f"镜头 {shot}: 包含模板化表达 `{phrase}`，应改成具体画面/动作/镜头描述")
        blocks = extract_blocks(prompt)
        if not blocks:
            errors.append(f"镜头 {shot}: 未找到逐秒画面内容块")
            continue
        for index, block in enumerate(blocks, 1):
            for label in REQUIRED_LABELS:
                if label not in block:
                    errors.append(f"镜头 {shot} 分镜 {index}: 缺少 `{label}`")

    if errors:
        print("\\n".join(errors))
        return 1
    print(f"PASS: {path.name} 已通过逐秒 Prompt 细节校验")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
