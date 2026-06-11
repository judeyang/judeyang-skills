#!/usr/bin/env python3
"""Validate mandatory per-segment detail labels in an internal video Prompt workbook."""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from openpyxl import load_workbook


REQUIRED_LABELS = ["景别：", "机位：", "构图：", "运镜手法：", "画面内容：", "特效：", "声音：", "结束状态：", "衔接要求："]
REQUIRED_PROMPT_SECTIONS = ["【基础设定】", "【氛围与画质】", "【画面内容】"]
REQUIRED_BASE_LABELS = ["镜头任务：", "参考素材：", "人物/产品/道具/场景：", "连续性：", "声音：", "文字策略："]
REQUIRED_QUALITY_LABELS = ["风格核心：", "视觉基调：", "色彩与影调："]
DEPRECATED_VIDEO_SECTIONS = ["【运镜规则】", "【声音/台词】", "【负面要求】", "【关键帧调用】"]
FORBIDDEN_ALTERNATE_HEADINGS = ["【视频类型】", "【产品一致性", "【三段式结构", "【逐秒分镜】", "【声音要求】"]
PLACEHOLDERS = ["按本镜头需要选择", "根据主体动作采用", "按画面执行", "保持可识别", "适当", "根据实际情况"]
GENERIC_TEMPLATE_PHRASES = [
    "围绕本段动作组织画面",
    "关键人物、手部动作、道具和视线方向完整入画",
    "视线或运动方向预留空间",
    "视线和动作方向预留空间",
    "按9:16竖屏组织前景、中景、背景层次",
    "固定机位，仅保留轻微呼吸感",
    "主体位于画面视觉中心",
    "主体置于画面视觉中心",
    "前景、中景、背景层次清楚",
    "固定机位或缓慢跟随主体移动",
    "固定机位或缓慢推进",
    "只生成只保留",
    "锁定本镜头中段的动作、表情或视觉变化",
    "锁定中段动作、表情或视觉变化",
    "锁定第1个中段动作或表情状态",
    "锁定第2个中段动作或表情状态",
]
TEXT_FORBID_MARKERS = ["禁止生成字幕", "水印", "乱码", "伪Logo"]
NO_TEXT_MARKERS = ["无文字纯图片", "默认画面为无文字", "默认无文字"]
ALLOWED_TEXT_MARKERS = ["文字策略", "产品镜头仅保留官方素材", "剧情道具", "报纸", "信件", "招牌", "包装", "屏显"]
FULL_PATH_MARKERS = ["/Users/", "\\Users\\", "file://"]
MUSIC_PROMPT_MARKERS = ["音乐/音效：", "古筝", "琵琶", "背景乐", "合奏", "轮指", "变奏"]
MUSIC_BOUNDARY_MARKERS = ["不需要配乐", "不生成BGM", "不生成音乐/BGM", "不生成音乐", "不要配乐"]
KEYFRAME_FILE_RE = re.compile(
    r"镜头\d+_(?:首帧|中间关键帧[A-Z]?|结尾帧)(?:_[^\\s；;,，。]+)?_v\\d+(?:_[^\\s；;,，。]+)?\\.(?:png|jpg|jpeg|webp)"
)
SOLID_COLOR_KEYFRAME_FILE_RE = re.compile(r"镜头\d+_[^\s；;,，。]*(?:纯黑|纯白|纯色|黑场|白场)[^\s；;,，。]*\.(?:png|jpg|jpeg|webp)")
SOLID_COLOR_KEYFRAME_PHRASES = ["纯黑关键帧", "纯白关键帧", "纯色关键帧", "黑场关键帧", "白场关键帧", "纯黑画面作为关键帧", "纯白画面作为关键帧"]
PRODUCT_RULE_MARKERS = ["机身比例", "Logo位置", "Logo错误", "出风口", "导风板", "屏显", "官方产品素材", "产品结构变形"]
PRODUCT_REF_MARKERS = ["产品@", "CAP", "卡萨帝空调", "空调", "170°", "170度", "气流", "送风", "导风板", "出风口", "屏显"]
QUALITY_BOUNDARY_RE = re.compile(r"质量边界：(?P<body>.*?)(?:\n|$)")
STYLE_IN_SHOT_SIZE_MARKERS = ["电影感", "史诗感", "高级感", "质感", "氛围"]
CAMERA_IN_COMPOSITION_MARKERS = ["无人机", "俯拍", "仰拍", "平视", "低机位", "高机位", "贴地", "正面30度", "侧面30度"]
PHOTOREALISM_TYPO_MARKERS = ["Photirealism", "Photorealisim", "Photorealstic"]
NEXT_SHOT_HANDOFF_MARKERS = ["切到", "切入", "切吴用", "切晁盖", "切杨志", "转到", "转入", "下一镜", "下个镜头", "接吴用", "接晁盖", "接杨志"]
HANDOFF_BOUNDARY_MARKERS = ["不生成下一镜", "不生成其他镜头", "不生成其他角色新动作", "不生成额外剧情", "本镜头停在"]
SOURCE_TERM_SPLIT_RE = re.compile(r"[，。；、：：“”！!？?\s+\-—\n（）()]+")
SOURCE_WEAK_TERMS = {
    "画面",
    "台词",
    "表演",
    "音效",
    "音乐",
    "镜头",
    "声音",
    "无",
    "在",
    "中",
    "的",
    "了",
    "和",
    "等人",
    "边",
    "继续",
    "已经",
    "开始",
    "后期",
    "合成",
    "执行",
    "内容",
    "车上有",
    "继续讲",
    "很开心",
}
SOURCE_KEYWORDS = [
    "黄泥冈",
    "烈日",
    "杨志",
    "押送",
    "生辰纲",
    "旗帜",
    "擦汗",
    "暗处",
    "草丛",
    "晁盖",
    "吴用",
    "埋伏",
    "怒视",
    "压低声音",
    "手机",
    "展示",
    "屏幕",
    "产品",
    "特效",
    "不屑",
    "抱臂",
    "算账",
    "占位",
    "省电",
    "数据",
    "卖点",
    "眼睛",
    "微亮",
    "嘴硬",
    "点头",
    "怀疑",
    "心动",
    "挠头",
    "比划",
    "幻想",
    "梁山",
    "好汉",
    "吃肉",
    "喝酒",
    "推眼镜",
    "愣",
    "拍腿",
    "拍大腿",
    "转身",
    "挥手",
    "离开",
    "齐声",
    "口号",
    "黑屏转场",
    "原地",
    "懵",
    "远去",
    "风",
    "落叶",
    "树林深处",
    "喊话",
    "独自",
    "声嘶力竭",
    "蝉鸣",
    "车轮",
    "紧张鼓点",
    "华丽出场",
    "哼",
    "林间风声",
    "归零",
    "轻快",
    "电子",
    "树木沙沙声",
    "树叶沙沙声",
    "锣鼓",
    "回声",
    "欢快",
    "乌鸦",
    "渐弱",
]


def find_header(ws, name: str) -> tuple[int, int]:
    for row in range(1, min(ws.max_row, 30) + 1):
        for col in range(1, ws.max_column + 1):
            if ws.cell(row, col).value == name:
                return row, col
    raise RuntimeError(f"未找到列：{name}")


def find_any_header(ws, names: list[str]) -> tuple[int, int]:
    errors = []
    for name in names:
        try:
            return find_header(ws, name)
        except RuntimeError as exc:
            errors.append(str(exc))
    raise RuntimeError(" / ".join(errors))


def maybe_find_header(ws, names: list[str]) -> int | None:
    for name in names:
        try:
            return find_header(ws, name)[1]
        except RuntimeError:
            continue
    return None


def find_prompt_sheet(wb):
    for ws in wb.worksheets:
        try:
            find_header(ws, "视频内容Prompt（含声音/负面）")
            find_header(ws, "镜号")
            return ws
        except RuntimeError:
            continue
    raise RuntimeError("未找到包含 `镜号` 和 `视频内容Prompt（含声音/负面）` 的执行表。")


def extract_blocks(prompt: str) -> list[str]:
    if "【画面内容】" not in prompt:
        return []
    content = prompt.rsplit("【画面内容】", 1)[1]
    match = re.search(r"\n【[^】]+】", content)
    if match:
        content = content[: match.start()]
    return [block.strip() for block in re.split(r"(?=分镜\d+：|(?=\d{2}:\d{2}(?:\.\d+)?-\d{2}:\d{2}(?:\.\d+)?\s*·))", content) if block.strip()]


def extract_base_section(prompt: str) -> str:
    if "【基础设定】" not in prompt:
        return ""
    section = prompt.split("【基础设定】", 1)[1]
    match = re.search(r"\n【[^】]+】", section)
    if match:
        section = section[: match.start()]
    return section.strip()


def keyframe_files(text: str) -> list[str]:
    seen: set[str] = set()
    files: list[str] = []
    for filename in KEYFRAME_FILE_RE.findall(text):
        if filename not in seen:
            seen.add(filename)
            files.append(filename)
    return files


def has_valid_text_strategy(prompt: str) -> bool:
    has_forbidden_baseline = all(marker in prompt for marker in TEXT_FORBID_MARKERS)
    has_default_no_text = any(marker in prompt for marker in NO_TEXT_MARKERS)
    has_allowed_exception = any(marker in prompt for marker in ALLOWED_TEXT_MARKERS)
    return has_forbidden_baseline and (has_default_no_text or has_allowed_exception)


def quality_negative_count(prompt: str) -> int:
    match = QUALITY_BOUNDARY_RE.search(prompt)
    if not match:
        return 0
    body = match.group("body")
    return body.count("避免") + body.count("禁止") + body.count("不要") + body.count("杜绝")


def label_value(block: str, label: str) -> str:
    pattern = rf"{re.escape(label)}(?P<value>.*?)(?:\n\S+：|\Z)"
    match = re.search(pattern, block, re.S)
    return match.group("value").strip() if match else ""


def normalize_source_text(value: str) -> str:
    text = str(value or "")
    text = re.sub(r"\s+", "", text)
    return (
        text.replace("“", '"')
        .replace("”", '"')
        .replace("‘", "'")
        .replace("’", "'")
        .replace("：", ":")
    )


def dialogue_fragments(dialogue: str) -> list[str]:
    text = str(dialogue or "").strip()
    if not text or text == "—":
        return []
    quoted = re.findall(r"[：:]“(.+?)”", text)
    if quoted:
        return quoted
    return [text]


def source_terms(text: str) -> list[str]:
    seen: set[str] = set()
    terms: list[str] = []
    for part in SOURCE_TERM_SPLIT_RE.split(str(text or "")):
        term = part.strip()
        if len(term) < 2:
            continue
        if term in SOURCE_WEAK_TERMS:
            continue
        if term.startswith("待") or term.endswith("待提供"):
            continue
        if term not in seen:
            seen.add(term)
            terms.append(term)
    return terms


def source_term_present(term: str, prompt: str) -> bool:
    normalized_term = normalize_source_text(term)
    normalized_prompt = normalize_source_text(prompt)
    if normalized_term in normalized_prompt:
        return True
    # A few production-safe equivalences used in prompt writing.
    aliases = {
        "树林风吹过": ["林间风声", "风声"],
        "树木沙沙声": ["树叶沙沙声", "树叶", "沙沙声"],
        "渐暗渐明黑屏转场": ["黑屏转场", "黑场转入"],
        "一脸懵逼": ["一脸懵", "懵"],
        "声嘶力竭": ["拖长声音", "喊话"],
        "轻快电子音": ["轻快电子", "轻快电子点按声", "电子点按声"],
        "林间风声": ["风声", "树林风声", "林间风"],
        "回声欢快音乐收尾": ["回声", "欢快氛围", "收尾欢快"],
        "吴用凑过来": ["吴用", "凑近"],
        "特写手机屏幕": ["手机", "屏幕", "特写"],
    }
    if any(normalize_source_text(alias) in normalized_prompt for alias in aliases.get(term, [])):
        return True
    term_keywords = [keyword for keyword in SOURCE_KEYWORDS if keyword in term]
    if term_keywords:
        present_count = sum(1 for keyword in term_keywords if normalize_source_text(keyword) in normalized_prompt)
        required_count = min(2, len(term_keywords))
        return present_count >= required_count
    if len(normalized_term) <= 4:
        return normalized_term[:2] in normalized_prompt
    return False


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True)
    args = parser.parse_args()

    path = Path(args.input).expanduser()
    wb = load_workbook(path, data_only=True)
    ws = find_prompt_sheet(wb)
    header_row, prompt_col = find_header(ws, "视频内容Prompt（含声音/负面）")
    shot_col = find_header(ws, "镜号")[1]
    _, refs_col = find_any_header(ws, ["参考资产", "参考素材"])
    exec_note_col = find_header(ws, "生成前执行说明")[1]
    first_col = find_header(ws, "首帧Prompt")[1]
    middle_col = find_header(ws, "中间关键帧Prompt")[1]
    end_col = find_header(ws, "结尾帧Prompt")[1]
    picture_col = maybe_find_header(ws, ["确认后执行内容", "确认后执行画面", "执行版画面内容", "画面内容", "画面"])
    dialogue_col = maybe_find_header(ws, ["台词/旁白", "确认后台词/旁白", "执行版台词/旁白", "台词"])
    sfx_col = maybe_find_header(ws, ["音效/音乐", "音效", "音乐"])
    errors: list[str] = []

    for row in range(header_row + 1, ws.max_row + 1):
        shot = ws.cell(row, shot_col).value
        if shot in (None, "", "合计"):
            continue
        prompt = str(ws.cell(row, prompt_col).value or "")
        refs_text = str(ws.cell(row, refs_col).value or "")
        exec_note = str(ws.cell(row, exec_note_col).value or "")
        first_prompt = str(ws.cell(row, first_col).value or "")
        middle_prompt = str(ws.cell(row, middle_col).value or "")
        end_prompt = str(ws.cell(row, end_col).value or "")
        keyframe_prompt_text = "\n".join([first_prompt, middle_prompt, end_prompt])
        picture_text = str(ws.cell(row, picture_col).value or "") if picture_col else ""
        dialogue_text = str(ws.cell(row, dialogue_col).value or "") if dialogue_col else ""
        sfx_text = str(ws.cell(row, sfx_col).value or "") if sfx_col else ""
        for fragment in dialogue_fragments(dialogue_text):
            if normalize_source_text(fragment) not in normalize_source_text(prompt):
                errors.append(f"镜头 {shot}: `台词/旁白` 中的完整台词未出现在视频Prompt中：{fragment}")
        for term in source_terms(picture_text):
            if not source_term_present(term, prompt):
                errors.append(f"镜头 {shot}: `确认后执行内容` 要点未出现在视频Prompt中：{term}")
        for term in source_terms(sfx_text):
            if not source_term_present(term, prompt):
                errors.append(f"镜头 {shot}: `音效/音乐` 要点未出现在视频Prompt中：{term}")
        for phrase in SOLID_COLOR_KEYFRAME_PHRASES:
            if phrase in keyframe_prompt_text or phrase in exec_note:
                errors.append(f"镜头 {shot}: 纯黑/纯白/纯色占位画面不能作为 AI 关键帧生成，请改为视频或后期转场说明")
        for filename in SOLID_COLOR_KEYFRAME_FILE_RE.findall(exec_note + "\n" + prompt):
            errors.append(f"镜头 {shot}: `{filename}` 是纯色/黑场占位帧，不应作为 AI 关键帧接入")
        if first_prompt.startswith("无需单独生成"):
            errors.append(f"镜头 {shot}: 首帧Prompt 不能标记为无需单独生成")
        first_reuses_previous = first_prompt.startswith("复用上一镜头结尾帧")
        needs_middle = "中间关键帧=生成" in exec_note or "中间关键帧：生成" in exec_note
        needs_end = "结尾帧=生成" in exec_note or "结尾帧：生成" in exec_note
        if first_reuses_previous and "复用上一镜头结尾帧" not in exec_note:
            errors.append(f"镜头 {shot}: 首帧复用上一镜头结尾帧，但生成前执行说明没有同步说明复用关系")
        if not needs_middle and middle_prompt.strip():
            errors.append(f"镜头 {shot}: 生成前执行说明未要求中间关键帧，但中间关键帧Prompt未留空")
        if needs_middle and not middle_prompt.strip():
            errors.append(f"镜头 {shot}: 生成前执行说明要求中间关键帧，但中间关键帧Prompt为空")
        if needs_middle and middle_prompt.startswith("无需单独生成"):
            errors.append(f"镜头 {shot}: 生成前执行说明要求中间关键帧，但中间关键帧Prompt被标记为无需生成")
        if not needs_end and end_prompt.strip():
            errors.append(f"镜头 {shot}: 生成前执行说明未要求结尾帧，但结尾帧Prompt未留空")
        if needs_end and not end_prompt.strip():
            errors.append(f"镜头 {shot}: 生成前执行说明要求结尾帧，但结尾帧Prompt为空")
        if needs_end and end_prompt.startswith("无需单独生成"):
            errors.append(f"镜头 {shot}: 生成前执行说明要求结尾帧，但结尾帧Prompt被标记为无需生成")
        has_product_ref = any(marker in (prompt + "\n" + refs_text + "\n" + exec_note) for marker in PRODUCT_REF_MARKERS)
        if not has_product_ref:
            for marker in PRODUCT_RULE_MARKERS:
                if marker in prompt:
                    errors.append(f"镜头 {shot}: 非产品镜头包含产品专用规则 `{marker}`，请按本镜头相关性裁剪")
        for section in REQUIRED_PROMPT_SECTIONS:
            if section not in prompt:
                errors.append(f"镜头 {shot}: 缺少三段式 Prompt 结构 `{section}`")
        for label in REQUIRED_QUALITY_LABELS:
            if label not in prompt:
                errors.append(f"镜头 {shot}: 【氛围与画质】缺少 `{label}`")
        base_section = extract_base_section(prompt)
        for label in REQUIRED_BASE_LABELS:
            if label not in base_section:
                errors.append(f"镜头 {shot}: 【基础设定】缺少 `{label}`")
        for section in DEPRECATED_VIDEO_SECTIONS:
            if section in prompt:
                errors.append(f"镜头 {shot}: 视频Prompt仍包含旧结构 `{section}`，请合并进三段式结构")
        for heading in FORBIDDEN_ALTERNATE_HEADINGS:
            if heading in prompt:
                errors.append(f"镜头 {shot}: 视频Prompt包含替代标题 `{heading}`，必须使用【基础设定】/【氛围与画质】/【画面内容】")
        if re.search(r"(^|\n)核心主题：", prompt):
            errors.append(f"镜头 {shot}: 视频Prompt仍包含旧结构 `核心主题：`，请改为【基础设定】里的`镜头任务：`")
        if "关键帧出图" not in exec_note or "视频生成前上传" not in exec_note:
            errors.append(f"镜头 {shot}: 生成前执行说明缺少关键帧出图或视频生成前上传说明")
        if "关键帧视觉复核" not in exec_note:
            errors.append(f"镜头 {shot}: 生成前执行说明缺少关键帧视觉复核要求")
        exec_files = keyframe_files(exec_note)
        prompt_files = keyframe_files(prompt)
        for filename in exec_files:
            if filename not in prompt_files:
                errors.append(f"镜头 {shot}: 生成前执行说明中的关键帧 `{filename}` 未出现在视频Prompt的参考素材说明中")
        if not needs_middle and "中间关键帧" not in base_section:
            errors.append(f"镜头 {shot}: 生成前执行说明跳过中间关键帧，但【基础设定】未说明中间关键帧不使用")
        if not needs_end and "结尾帧" not in base_section:
            errors.append(f"镜头 {shot}: 生成前执行说明跳过结尾帧，但【基础设定】未说明结尾帧不使用")
        if "引用资产：" in prompt:
            errors.append(f"镜头 {shot}: 视频Prompt包含散乱引用资产执行说明，应移到生成前执行说明并在画面描述中使用 inline asset tags")
        if not any(marker in prompt for marker in MUSIC_BOUNDARY_MARKERS):
            errors.append(f"镜头 {shot}: 视频Prompt缺少音乐/BGM不生成或后期单独配的声音边界")
        for marker in MUSIC_PROMPT_MARKERS:
            if marker in prompt:
                errors.append(f"镜头 {shot}: 视频Prompt包含音乐生成相关表达 `{marker}`，视频工具只生成音效/同期声")
        for marker in FULL_PATH_MARKERS:
            if marker in prompt or marker in refs_text:
                errors.append(f"镜头 {shot}: 参考素材或 Prompt 中包含完整本地路径 `{marker}`，应改为文件名")
        for marker in PHOTOREALISM_TYPO_MARKERS:
            if marker in prompt:
                errors.append(f"镜头 {shot}: `{marker}` 拼写错误，应写为 `Photorealism`")
        if not has_valid_text_strategy(prompt):
            errors.append(f"镜头 {shot}: 缺少有效文字策略；默认应禁字幕/水印/乱码/伪Logo，产品或剧情道具文字必须明确例外")
        negative_count = quality_negative_count(prompt)
        if negative_count > 8:
            errors.append(f"镜头 {shot}: 质量边界负面项过多（{negative_count}项），请裁剪到本镜头相关的3-6项，最多8项")
        for placeholder in PLACEHOLDERS:
            if placeholder in prompt:
                errors.append(f"镜头 {shot}: 包含占位表达 `{placeholder}`")
        for phrase in GENERIC_TEMPLATE_PHRASES:
            if phrase in prompt:
                errors.append(f"镜头 {shot}: 包含模板化表达 `{phrase}`，应改成具体画面/动作/镜头描述")
        blocks = extract_blocks(prompt)
        if not blocks:
            errors.append(f"镜头 {shot}: 未找到逐秒画面内容块")
            continue
        for index, block in enumerate(blocks, 1):
            for label in REQUIRED_LABELS:
                if label not in block:
                    errors.append(f"镜头 {shot} 分镜 {index}: 缺少 `{label}`")
            shot_size = label_value(block, "景别：")
            composition = label_value(block, "构图：")
            handoff = label_value(block, "衔接要求：")
            for marker in STYLE_IN_SHOT_SIZE_MARKERS:
                if marker in shot_size:
                    errors.append(f"镜头 {shot} 分镜 {index}: `景别` 中包含风格词 `{marker}`，请移到【氛围与画质】")
            for marker in CAMERA_IN_COMPOSITION_MARKERS:
                if marker in composition:
                    errors.append(f"镜头 {shot} 分镜 {index}: `构图` 中包含机位词 `{marker}`，请移到`机位`")
            for marker in NEXT_SHOT_HANDOFF_MARKERS:
                if marker in handoff:
                    errors.append(f"镜头 {shot} 分镜 {index}: `衔接要求` 不应描述下一镜具体内容 `{marker}`；请只写本镜头结束边界")
            if handoff and not any(marker in handoff for marker in HANDOFF_BOUNDARY_MARKERS):
                errors.append(f"镜头 {shot} 分镜 {index}: `衔接要求` 应作为结束边界使用，建议写 `本镜头停在上述结束状态；不生成下一镜、其他角色新动作或额外剧情。`")

    if errors:
        print("\n".join(errors))
        return 1
    print(f"PASS: {path.name} 已通过逐秒 Prompt 细节校验")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
