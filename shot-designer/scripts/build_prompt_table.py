#!/usr/bin/env python3
"""Build an internal execution workbook scaffold.

This script intentionally does not generate any prompt content. Prompt cells are
left blank for the AI director pass after the full script has been read.
"""

from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

CENTER_HEADERS = {
    "序号",
    "镜号",
    "原始镜号",
    "时长",
    "执行时长(s)",
    "景别",
    "制作状态",
    "输出文件名",
}

COMPACT_TEXT_HEADERS = {
    "确认后台词/旁白",
    "台词/旁白",
    "音效/音乐",
    "样片建议",
}


def body_alignment(header_value):
    header = str(header_value or "").strip()
    if header in CENTER_HEADERS or header.endswith("编号"):
        return Alignment(horizontal="center", vertical="top", wrap_text=True)
    if header in COMPACT_TEXT_HEADERS:
        return Alignment(horizontal="left", vertical="top", wrap_text=True)
    return Alignment(horizontal="left", vertical="top", wrap_text=True)


IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp"}


def normalize(value) -> str:
    return str(value or "").strip().replace("\n", " / ")


def find_project_root(path: Path) -> Path | None:
    for parent in [path.parent, *path.parents]:
        if (
            parent.name.endswith("_项目交付文件夹")
            or (parent / "03_设定资产").exists()
            or (parent / "04_设定资产").exists()
            or (parent / "03_设定图输出").exists()
        ):
            return parent
    return None


def find_header_row(ws) -> tuple[int, dict[str, int]] | None:
    canonical = {
        "shot": ["镜号", "执行镜号"],
        "orig": ["原始镜号", "对应原始镜号"],
        "orig_text": ["原始脚本内容"],
        "picture": ["确认后执行内容", "确认后执行画面", "执行版画面内容", "画面内容", "画面"],
        "dialogue": ["确认后台词/旁白", "执行版台词/旁白", "台词/旁白", "台词"],
        "duration": ["时长", "执行时长"],
        "sfx": ["音效/音乐", "音效", "音乐"],
        "note": ["改动说明", "执行备注", "备注"],
    }
    for row_idx in range(1, min(ws.max_row, 30) + 1):
        values = [normalize(ws.cell(row_idx, col).value) for col in range(1, ws.max_column + 1)]
        mapping: dict[str, int] = {}
        for col_idx, value in enumerate(values, 1):
            compact = value.replace(" ", "")
            for key, names in canonical.items():
                if key not in mapping and any(name in compact for name in names):
                    mapping[key] = col_idx
        if "shot" in mapping and "picture" in mapping:
            return row_idx, mapping
    return None


def read_execution_rows(path: Path) -> list[dict[str, str]]:
    wb = load_workbook(path, data_only=True)
    candidates = [wb["执行版脚本"]] if "执行版脚本" in wb.sheetnames else []
    candidates.extend(ws for ws in wb.worksheets if ws not in candidates)
    ws = None
    found = None
    for candidate in candidates:
        found = find_header_row(candidate)
        if found:
            ws = candidate
            break
    if not found:
        raise RuntimeError("未找到执行脚本表头，需要至少包含 镜号 和 确认后执行内容/确认后执行画面。")
    assert ws is not None
    header_row, mapping = found
    rows: list[dict[str, str]] = []
    for row_idx in range(header_row + 1, ws.max_row + 1):
        shot = normalize(ws.cell(row_idx, mapping["shot"]).value)
        if not shot:
            continue
        if shot in {"合计", "总计"}:
            rows.append({"total": "1", "duration": normalize(ws.cell(row_idx, mapping.get("duration", 0)).value) if mapping.get("duration") else ""})
            continue
        rows.append(
            {
                "shot": shot,
                "orig": normalize(ws.cell(row_idx, mapping.get("orig", 0)).value) if mapping.get("orig") else "",
                "orig_text": normalize(ws.cell(row_idx, mapping.get("orig_text", 0)).value) if mapping.get("orig_text") else "",
                "picture": normalize(ws.cell(row_idx, mapping["picture"]).value),
                "dialogue": normalize(ws.cell(row_idx, mapping.get("dialogue", 0)).value) if mapping.get("dialogue") else "",
                "duration": normalize(ws.cell(row_idx, mapping.get("duration", 0)).value) if mapping.get("duration") else "",
                "sfx": normalize(ws.cell(row_idx, mapping.get("sfx", 0)).value) if mapping.get("sfx") else "",
                "note": normalize(ws.cell(row_idx, mapping.get("note", 0)).value) if mapping.get("note") else "",
            }
        )
    return rows


def scan_assets(project_dir: Path | None) -> dict[str, Path]:
    if not project_dir:
        return {}
    root = project_dir / "03_设定资产"
    if not root.exists():
        root = project_dir / "04_设定资产"
    if not root.exists():
        root = project_dir / "03_设定图输出"
    if not root.exists():
        return {}
    assets: dict[str, Path] = {}
    for path in root.rglob("*"):
        if path.is_file() and path.suffix.lower() in IMAGE_EXTENSIONS:
            assets[path.stem] = path
    return assets


def asset_display(path: Path) -> str:
    return path.name


ASSET_ALIASES = {
    "贾母": ["贾母"],
    "林黛玉": ["林黛玉", "黛玉"],
    "薛宝钗": ["薛宝钗", "宝钗"],
    "贾宝玉": ["贾宝玉", "宝玉"],
    "王熙凤": ["王熙凤", "凤姐"],
    "刘姥姥": ["刘姥姥"],
    "群像": ["群像", "众人", "宴会", "宴席"],
    "宴会厅": ["宴会厅", "大观园", "宴席"],
    "C位座次": ["C位", "空座", "座次"],
    "道具": ["团扇", "绢帕", "器物", "杯盏", "道具"],
    "气流": ["气流", "送风", "170度", "环抱风"],
    "孙悟空": ["孙悟空", "悟空"],
    "唐三藏": ["唐三藏", "唐僧", "师父"],
    "猪八戒": ["猪八戒", "八戒", "老猪"],
    "沙僧": ["沙僧"],
    "白龙马": ["白龙马", "马"],
    "黄风怪": ["黄风怪"],
    "灵吉菩萨": ["灵吉菩萨", "菩萨"],
    "黄风岭": ["黄风岭", "黄沙", "风沙", "沙丘"],
    "花果山": ["花果山", "雨林", "瀑布", "彩虹"],
    "海尔空调": ["海尔", "空调", "产品", "水洗空气"],
    "卡萨帝空调": ["卡萨帝", "空调", "产品", "柜机", "揽光"],
    "法器": ["净风宝瓶", "定风珠", "法器", "神器"],
}


def matching_asset_paths(label: str, keywords: list[str], assets: dict[str, Path]) -> list[Path]:
    character_labels = {"贾母", "林黛玉", "薛宝钗", "贾宝玉", "王熙凤", "刘姥姥", "孙悟空", "唐三藏", "猪八戒", "沙僧", "白龙马", "黄风怪", "灵吉菩萨"}
    scene_labels = {"宴会厅", "C位座次", "黄风岭", "花果山"}
    matches: list[tuple[int, Path]] = []
    for name, path in assets.items():
        searchable = f"{name} {path.name}"
        if label in searchable or any(k in searchable for k in keywords):
            score = 5
            if label in searchable:
                score = 2
            if label in character_labels and f"角色_{label}" in searchable:
                score = 0
            if label in scene_labels and ("场景_" in searchable or label in searchable):
                score = min(score, 1)
            if label == "群像" and "群像" in searchable:
                score = 0
            if label == "道具" and "道具" in searchable:
                score = 0
            matches.append((score, path))
    return [path for _, path in sorted(matches, key=lambda item: (item[0], len(item[1].name), item[1].name))]


def references_for(row: dict[str, str], assets: dict[str, Path]) -> str:
    dialogue = row.get("dialogue", "")
    if "标题字" in dialogue:
        dialogue = ""
    # Use confirmed execution content for visual references. Original script text may
    # contain titles or replaced ideas that are no longer visible in the shot.
    text = f"{row.get('picture','')} {dialogue}"
    lines: list[str] = []
    for label, keywords in ASSET_ALIASES.items():
        if not any(k in text for k in keywords):
            continue
        matches = matching_asset_paths(label, keywords, assets)
        if matches:
            lines.append(f"{label}: {asset_display(matches[0])}")
        else:
            lines.append(f"{label}: 待补充")
    return "\n".join(lines) if lines else "按已确认设定和文字方向执行"


def estimate_row_height(row, header_row: int) -> float:
    if row[0].row == header_row:
        return 34
    max_chars = 0
    max_lines = 1
    for cell in row:
        text = str(cell.value or "")
        max_chars = max(max_chars, len(text))
        max_lines = max(max_lines, text.count("\n") + 1)
    # Keep rows scannable. Long Prompt cells remain wrapped and inspectable in the formula bar.
    by_length = 54 + min(max_chars / 95, 6) * 16
    by_lines = min(max_lines, 8) * 15
    return max(68, min(168, by_length, max(72, by_lines)))


def style_project_info_sheet(ws, title: str = "项目说明"):
    title_fill = PatternFill("solid", fgColor="1F4E78")
    label_fill = PatternFill("solid", fgColor="EAF3F8")
    thin = Side(style="thin", color="D9E2F3")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.sheet_view.zoomScale = 90
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 96
    ws.freeze_panes = "A2"

    ws.insert_rows(1)
    ws["A1"] = title
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=2)
    ws["A1"].fill = title_fill
    ws["A1"].font = Font(name="Microsoft YaHei", size=13, bold=True, color="FFFFFF")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    for row in ws.iter_rows(min_row=2, max_col=2):
        for cell in row:
            cell.border = border
            cell.font = Font(name="Microsoft YaHei", size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        row[0].fill = label_fill
        row[0].font = Font(name="Microsoft YaHei", size=10, bold=True, color="1F4E78")
        row[0].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        row[1].alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        text = str(row[1].value or "")
        row_idx = row[0].row
        ws.row_dimensions[row_idx].height = max(28, min(96, 24 + text.count("\n") * 16 + len(text) / 90 * 14))


def style_sheet(ws, widths: list[int], header_row: int = 1):
    header_fill = PatternFill("solid", fgColor="1F4E78")
    subtle_fill = PatternFill("solid", fgColor="D9EAF7")
    thin = Side(style="thin", color="D9E2F3")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for cell in ws[header_row]:
        cell.fill = header_fill
        cell.font = Font(name="Microsoft YaHei", color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border

    for row in ws.iter_rows(min_row=header_row + 1):
        for cell in row:
            cell.font = Font(name="Microsoft YaHei", size=10)
            cell.alignment = body_alignment(ws.cell(header_row, cell.column).value)
            cell.border = border
        if row[0].value == "合计":
            for cell in row:
                cell.fill = subtle_fill
                cell.font = Font(name="Microsoft YaHei", size=10, bold=True)

    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = width
    for row_index in range(1, ws.max_row + 1):
        ws.row_dimensions[row_index].height = estimate_row_height(ws[row_index], header_row)
    ws.freeze_panes = f"A{header_row + 1}"
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.zoomScale = 80


def default_output_path(input_path: Path, project_dir: Path | None) -> Path:
    filename = f"{input_path.stem}_内部执行脚本与Prompt表_v01.xlsx"
    if project_dir:
        target_dir = project_dir / "05_内部制作执行"
        if not target_dir.exists():
            target_dir = project_dir / "06_内部制作执行"
        if not target_dir.exists():
            target_dir = project_dir / "04_分镜与视频生成"
        return target_dir / filename
    return input_path.with_name(filename)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True, help="Approved client confirmation or internal structure workbook")
    parser.add_argument("--output", help="Output prompt workbook path")
    parser.add_argument("--project-dir", help="Optional project root folder")
    args = parser.parse_args()

    input_path = Path(args.input).expanduser()
    project_dir = Path(args.project_dir).expanduser() if args.project_dir else find_project_root(input_path)
    output_path = Path(args.output).expanduser() if args.output else default_output_path(input_path, project_dir)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    rows = read_execution_rows(input_path)
    assets = scan_assets(project_dir)

    wb = Workbook()
    ws = wb.active
    ws.title = "执行脚本与Prompt表"
    info_ws = wb.create_sheet("项目说明")
    meta = [
        ["项目", project_dir.name.replace("_项目交付文件夹", "") if project_dir else input_path.stem],
        ["文件用途", "内部制作执行表结构：同一行查看原始脚本、确认后执行脚本和参考素材；Prompt列留空，等待AI导演逐镜头撰写。"],
        ["画面规格", "9:16竖屏"],
        ["Prompt标准", "脚本不生成Prompt。AI通读全片后先锁定连续性，再按三段式填写视频内容Prompt：【基础设定】、【氛围与画质】、【画面内容】；运镜、声音和负面边界合并进对应段落。"],
        ["执行口径", "分镜图不再提交用户确认；本脚本仅生成结构和参考素材，不进入首帧图或视频生成。"],
        ["版本", "v01"],
    ]
    for meta_row in meta:
        info_ws.append(meta_row)
    ws.append(["镜号", "原始镜号", "原始脚本内容", "确认后执行内容", "执行时长(s)", "台词/旁白", "音效/音乐", "参考素材", "首帧Prompt", "中间关键帧Prompt", "结尾帧Prompt", "生成前执行说明", "视频内容Prompt（含声音/负面）", "用户确认/修改依据", "样片建议", "制作状态", "输出文件名", "备注"])

    for row in rows:
        if row.get("total"):
            ws.append(["合计", "", "", "", row.get("duration"), "", "", "", "", "", "", "", "", "", "", "", "", ""])
            continue
        refs = references_for(row, assets)
        ws.append(
            [
                row.get("shot"),
                row.get("orig"),
                row.get("orig_text"),
                row.get("picture"),
                row.get("duration"),
                row.get("dialogue"),
                row.get("sfx"),
                refs,
                "",
                "",
                "",
                "",
                "",
                row.get("note"),
                "",
                "待AI逐镜头撰写Prompt",
                "",
                "脚本仅生成结构；Prompt列必须由AI通读全片后填写。",
            ]
        )

    style_sheet(ws, [7, 8, 36, 32, 10, 15, 16, 36, 50, 34, 34, 50, 78, 24, 18, 11, 18, 18], header_row=1)
    style_project_info_sheet(info_ws)
    wb.save(output_path)
    print(output_path)


if __name__ == "__main__":
    main()
